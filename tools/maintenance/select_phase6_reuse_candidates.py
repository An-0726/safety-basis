#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Select PHASE 6 candidates whose cited basis matches a current reviewed clause.

This is deliberately a selection/reporting tool, not a promotion tool.  It
reads the authoritative workbook and ``knowledge/``; it never changes a
hazard, link, review, manifest, SQLite database, private library, or release
bundle.  An exact clause match only means that a human reviewer has a
reusable starting point.  Link applicability and the current official source
still have to be checked before any promotion batch.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
KNOWLEDGE = ROOT / "knowledge"
DEFAULT_XLSX = Path(r"D:/Desktop/隐患库_1929条_新版口径全部整改完成_20260914.xlsx")
DEFAULT_AS_OF = "2026-09-16"

sys.path.insert(0, str(ROOT / "tools" / "v4"))
from canonical import content_hash  # noqa: E402
from release_gate_core import evaluate_release_gate  # noqa: E402


def load_dir(relative: str | Path) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for path in sorted((KNOWLEDGE / relative).glob("*.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        result[str(data.get("id") or data.get("entityId") or path.stem)] = data
    return result


def compact(value: Any) -> str:
    return re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]", "", str(value or "")).casefold()


def normalize_article_locator(value: Any) -> str:
    text = str(value or "").strip()
    text = re.sub(r"^第\s*", "", text)
    text = re.sub(r"\s*条$", "", text)
    return text.replace("（", "(").replace("）", ")").replace(" ", "")


def article_tokens(value: Any) -> set[str]:
    text = str(value or "")
    found = re.findall(
        r"第\s*([0-9]+(?:\.[0-9]+)*|[一二三四五六七八九十百千万零〇两]+)\s*条",
        text,
    )
    found.extend(re.findall(r"(?<![A-Za-z])([0-9]+(?:\.[0-9]+)+)", text))
    return {normalize_article_locator(item) for item in found}


def read_scope() -> dict[str, str]:
    result: dict[str, str] = {}
    path = ROOT / "docs" / "hazard-reconciliation.jsonl"
    if not path.exists():
        return result
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        item = json.loads(line)
        if item.get("recordType") == "target":
            result[str(item["hazardId"])] = str(item.get("classification") or "unknown")
    return result


def review_summary(
    entity: dict[str, Any],
    review: dict[str, Any] | None,
    evidence: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    if not review:
        return {"decision": "missing", "hashCurrent": False, "evidenceRefs": []}
    refs = [item for item in (review.get("evidenceRefs") or []) if isinstance(item, str)]
    return {
        "decision": review.get("decision") or "missing",
        "hashCurrent": review.get("reviewedContentHash") == content_hash(entity),
        "evidenceRefs": refs,
        "evidenceFoundCount": sum(item in evidence for item in refs),
        "checkedAt": review.get("checkedAt") or review.get("reviewedAt"),
        "reviewer": review.get("reviewer"),
    }


def find_candidates(xlsx: Path, as_of_text: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    as_of = date.fromisoformat(as_of_text)
    hazards = load_dir("hazards")
    hazard_reviews = load_dir("reviews/hazards")
    clauses = load_dir("clauses")
    clause_reviews = load_dir("reviews/clauses")
    law_versions = load_dir("law-versions")
    evidence = load_dir("evidence")
    scope = read_scope()
    proposed = {hid for hid, hazard in hazards.items() if hazard.get("lifecycle") == "proposed"}

    # Normalize articlePath so both “第5.7.4条” and “5.7.4” resolve to the
    # same locator.  This avoids the silent zero-match failure of older dry
    # runs while preserving the exact clause object and review evidence.
    clauses_by_version: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for clause in clauses.values():
        key = (str(clause.get("lawVersionId")), normalize_article_locator(clause.get("articlePath")))
        clauses_by_version.setdefault(key, []).append(clause)

    gate = evaluate_release_gate(KNOWLEDGE, as_of=as_of)
    workbook = load_workbook(xlsx, read_only=True, data_only=True)
    sheet = workbook["隐患明细_修订后"]
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(value or "") for value in rows[0]]
    indexes = {name: index for index, name in enumerate(headers)}
    required = {"隐患ID", "直接依据", "修订状态"}
    missing = required - set(indexes)
    if missing:
        raise ValueError(f"workbook missing columns: {sorted(missing)}")

    records: list[dict[str, Any]] = []
    for row in rows[1:]:
        hazard_id = str(row[indexes["隐患ID"]] or "").strip()
        if hazard_id not in proposed:
            continue
        basis = str(row[indexes["直接依据"]] or "").strip()
        matches: dict[str, dict[str, Any]] = {}
        for line in basis.splitlines():
            line_text = compact(line)
            for version_id, version in law_versions.items():
                document_number = compact(version.get("documentNumber"))
                official_name = compact(version.get("officialName"))
                if not (
                    (document_number and len(document_number) >= 5 and document_number in line_text)
                    or (official_name and len(official_name) >= 6 and official_name in line_text)
                ):
                    continue
                gate_status = gate.law_versions.get(version_id, {})
                if not gate_status.get("ok") or not gate_status.get("supports_current"):
                    continue
                for locator in article_tokens(line):
                    for clause in clauses_by_version.get((version_id, locator), []):
                        clause_gate = gate.clauses.get(clause["id"], {})
                        if not clause_gate.get("ok"):
                            continue
                        matches[clause["id"]] = clause

        if not matches:
            continue
        hazard = hazards[hazard_id]
        hazard_review = review_summary(hazard, hazard_reviews.get(hazard_id), evidence)
        hazard_review_ready = (
            hazard_review["decision"] == "verified" and hazard_review["hashCurrent"]
        )
        clause_rows = []
        for clause in sorted(matches.values(), key=lambda item: item["id"]):
            version = law_versions.get(clause.get("lawVersionId"), {})
            clause_review = review_summary(clause, clause_reviews.get(clause["id"]), evidence)
            clause_rows.append(
                {
                    "clauseId": clause["id"],
                    "articlePath": clause.get("articlePath"),
                    "lawVersionId": clause.get("lawVersionId"),
                    "lawVersionName": version.get("officialName"),
                    "documentNumber": version.get("documentNumber"),
                    "validityStatus": version.get("validityStatus"),
                    "effectiveDate": version.get("effectiveDate"),
                    "sourceUrl": clause.get("sourceUrl") or version.get("sourceUrl"),
                    "clauseReview": clause_review,
                    "quotePresent": bool(clause.get("quote")),
                }
            )

        target_scope = scope.get(hazard_id)
        if not target_scope and hazard.get("proposalStatus") == "knowledge_extra_non_target_pending_scope_review":
            target_scope = "knowledge-extra-proposed-non-target"
        target_scope = target_scope or "unmapped"
        records.append(
            {
                "recordType": "reuseCandidate",
                "hazardId": hazard_id,
                "title": hazard.get("title"),
                "proposalStatus": hazard.get("proposalStatus"),
                "sourceRow": hazard.get("sourceRow"),
                "targetScope": target_scope,
                "basisLines": [line for line in basis.splitlines() if line.strip()],
                "hazardReview": hazard_review,
                "exactCurrentClauseCount": len(clause_rows),
                "exactCurrentClauses": clause_rows,
                "selectionClass": (
                    "manual_batch_1_hazard_review_verified"
                    if hazard_review_ready
                    else "manual_batch_2_hazard_review_missing_or_stale"
                ),
                "safeToPromoteNow": False,
                "nextAction": (
                    "复核 hazard 对象与条款适用性，并建立新的 direct/fallback link；不得直接沿用旧摘要转正。"
                    if hazard_review_ready
                    else "先补做当前 hash 的 hazard review，再复核条款适用性并建立关联；保持 proposed。"
                ),
            }
        )

    class_counts = Counter(item["selectionClass"] for item in records)
    metadata = {
        "recordType": "metadata",
        "schemaVersion": 1,
        "asOf": as_of_text,
        "purpose": "PHASE 6 manual selection aid; exact current reviewed clause reuse only",
        "source": {
            "workbook": str(xlsx).replace("\\", "/") if xlsx.is_absolute() else str(xlsx),
            "basisColumn": "直接依据",
            "matchingRule": "same current-gated lawVersion document/name + normalized article locator",
            "candidateLifecycle": "proposed",
        },
        "summary": {
            "matchedCandidates": len(records),
            "selectionClassCounts": dict(sorted(class_counts.items())),
            "targetScopeCounts": dict(sorted(Counter(item["targetScope"] for item in records).items())),
            "proposalStatusCounts": dict(sorted(Counter(item["proposalStatus"] for item in records).items())),
            "exactClauseCount": sum(item["exactCurrentClauseCount"] for item in records),
            "safeToPromoteNow": 0,
        },
        "integrity": {
            "allCandidatesProposed": all(item["hazardId"] in proposed for item in records),
            "uniqueHazardIds": len({item["hazardId"] for item in records}) == len(records),
            "allClausesCurrentGated": all(
                clause["validityStatus"] == "active"
                for item in records
                for clause in item["exactCurrentClauses"]
            ),
            "noPromotionPerformed": all(item["safeToPromoteNow"] is False for item in records),
        },
    }
    return metadata, records


def write_jsonl(path: Path, metadata: dict[str, Any], records: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        "".join(
            json.dumps(item, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n"
            for item in [metadata] + records
        ),
        encoding="utf-8",
    )


def table(headers: list[str], rows: list[list[str]]) -> str:
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in rows)
    return "\n".join(lines)


def write_markdown(path: Path, metadata: dict[str, Any], records: list[dict[str, Any]], as_of: str) -> None:
    summary = metadata["summary"]
    families = Counter()
    for record in records:
        for clause in record["exactCurrentClauses"]:
            families[clause["lawVersionName"] or clause["lawVersionId"]] += 1
    rows = [[key, str(value)] for key, value in sorted(families.items(), key=lambda item: (-item[1], item[0]))]
    lines = [
        "# PHASE 6 可复用条款候选批",
        "",
        f"> 核验基准日：{as_of}。这是人工核验选择清单，不是转正清单。所有候选仍为 `proposed`。",
        "> 匹配只说明：修订表直接依据中的法规/标准和条款定位，能映射到当前 Gate 已通过的 knowledge 条款；仍需重新核对对象、适用性和官方来源。",
        "",
        "## 选择结果",
        "",
        f"- 精确匹配候选：**{summary['matchedCandidates']}** 条；可作为第一人工批（hazard review 当前 hash 已验证）：**{summary['selectionClassCounts'].get('manual_batch_1_hazard_review_verified', 0)}** 条。",
        f"- 需要先补 hazard review 的第二批：**{summary['selectionClassCounts'].get('manual_batch_2_hazard_review_missing_or_stale', 0)}** 条。",
        f"- 匹配到当前条款对象：**{summary['exactClauseCount']}** 个；没有任何记录被标为可直接转正。",
        "",
        "## 按法规版本的可复用条款数",
        "",
        table(["法规/标准版本", "条款匹配数"], rows),
        "",
        "## 第一人工批执行门槛",
        "",
        "1. 逐条打开 `sourceUrl` 或合法保存的原始文件，重新确认现行版本、实施日期和条款逐字原文。",
        "2. 重新比较 hazard 的对象、场景、整改措施与条款义务；旧 hazard review 只作起点，不替代本批适用性审核。",
        "3. 新建 direct/fallback link 和当前上下文 hash；完整 Gate 通过前保持 `proposed`。",
        "4. 条款适用范围、现行性或证据链有疑问时标 `待核`，不因已有匹配强行转正。",
        "",
        "机器明细见 [`docs/phase6-reuse-candidates.jsonl`](phase6-reuse-candidates.jsonl)。",
        "",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--as-of", default=DEFAULT_AS_OF)
    parser.add_argument("--jsonl", type=Path, default=ROOT / "docs" / "phase6-reuse-candidates.jsonl")
    parser.add_argument("--markdown", type=Path, default=ROOT / "docs" / "PHASE6_REUSE_CANDIDATES.md")
    args = parser.parse_args()
    metadata, records = find_candidates(args.xlsx, args.as_of)
    write_jsonl(args.jsonl, metadata, records)
    write_markdown(args.markdown, metadata, records, args.as_of)
    print(json.dumps(metadata["summary"], ensure_ascii=False, sort_keys=True))
    return 0 if metadata["integrity"]["allCandidatesProposed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())

# -*- coding: utf-8 -*-
"""434 条 proposed 候选全量聚合与深度审查脚本。"""
from __future__ import annotations
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[2]
KNOW = ROOT / "knowledge"
EXCEL_PATH = Path(r"D:\Desktop\隐患库_1929条_新版口径全部整改完成_20260914.xlsx")

def load_json_dir(rel: str) -> dict[str, dict]:
    out = {}
    for p in (KNOW / rel).glob("*.json"):
        with open(p, "r", encoding="utf-8") as f:
            d = json.load(f)
        key = d.get("id") or d.get("entityId") or p.stem
        out[key] = d
    return out

def main():
    hazards = load_json_dir("hazards")
    clauses = load_json_dir("clauses")
    links = load_json_dir("links")
    lvs = load_json_dir("law-versions")
    cr = load_json_dir("reviews/clauses")

    proposed_ids = sorted(
        h["id"] for h in hazards.values() if h.get("lifecycle") == "proposed"
    )
    print(f"Total proposed hazards: {len(proposed_ids)}")

    # 1. 加载 Excel 1929
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sheet = wb["隐患明细_修订后"]
    headers = [cell.value for cell in sheet[1]]
    id_idx = headers.index("隐患ID")
    excel_rows = {}
    for r in range(2, sheet.max_row + 1):
        vals = [sheet.cell(r, c).value for c in range(1, len(headers) + 1)]
        hid = str(vals[id_idx]).strip() if vals[id_idx] else ""
        if hid:
            excel_rows[hid] = dict(zip(headers, vals))

    # 2. 建立知识库现行条款索引
    # active clauses with verified review
    active_clauses = {}
    for cid, c in clauses.items():
        if c.get("lifecycle") != "active":
            continue
        lvid = c.get("lawVersionId")
        lv = lvs.get(lvid, {})
        if lv.get("validityStatus") != "active":
            continue
        r = cr.get(cid, {})
        if r.get("decision") != "verified":
            continue
        active_clauses[cid] = {
            "clause": c,
            "lawVersion": lv,
            "review": r
        }
    print(f"Total verified active clauses under active law versions: {len(active_clauses)}")

    # 3. 聚合每条 proposed 的全部信息
    audit_records = []
    for hid in proposed_ids:
        h = hazards[hid]
        ex = excel_rows.get(hid, {})
        
        # 提取依据线索
        ex_basis = str(ex.get("直接依据") or "").strip()
        ex_key_points = str(ex.get("依据条款要点（修订，非原文摘录）") or "").strip()
        ex_backup = str(ex.get("补充/兜底依据") or "").strip()
        ex_revision_note = str(ex.get("修订说明") or "").strip()
        h_note = str(h.get("note") or "").strip()
        h_title = str(h.get("title") or "").strip()
        h_desc = str(h.get("description") or "").strip()
        
        # 提取标准号
        combined_text = f"{ex_basis} {ex_key_points} {ex_backup} {ex_revision_note} {h_note} {h_title}"
        
        audit_records.append({
            "hazardId": hid,
            "title": h_title,
            "category": h.get("category"),
            "proposalStatus": h.get("proposalStatus"),
            "sourceRow": h.get("sourceRow"),
            "ex_basis": ex_basis,
            "ex_key_points": ex_key_points,
            "ex_backup": ex_backup,
            "ex_revision_note": ex_revision_note,
            "h_note": h_note,
            "combined_text": combined_text,
            "hazard": h,
            "excel": ex
        })

    # 输出统计报告
    print(f"Aggregated {len(audit_records)} records successfully.")
    
    # 按照主要标准/法规族聚类
    clusters = defaultdict(list)
    for rec in audit_records:
        txt = rec["combined_text"]
        matched_standards = []
        if "13869" in txt:
            matched_standards.append("GB/T 13869")
        if "12801" in txt:
            matched_standards.append("GB 12801")
        if "47236" in txt:
            matched_standards.append("GB/T 47236")
        if "55037" in txt:
            matched_standards.append("GB 55037")
        if "55036" in txt:
            matched_standards.append("GB 55036")
        if "18597" in txt:
            matched_standards.append("GB 18597")
        if "12158" in txt:
            matched_standards.append("GB 12158")
        if "15577" in txt or "粉尘防爆安全规程" in txt or "粉尘防爆安全规定" in txt:
            matched_standards.append("粉尘防爆(GB15577/规定)")
        if "15603" in txt or "危化品仓库" in txt:
            matched_standards.append("GB 15603")
        if "安全生产法" in txt:
            matched_standards.append("安全生产法")
        if "消防法" in txt:
            matched_standards.append("消防法")
        if "特种设备安全法" in txt:
            matched_standards.append("特种设备安全法")
        if "46768" in txt or "有限空间" in txt:
            matched_standards.append("有限空间(GB46768/规定)")
        if "50187" in txt or "总平面设计" in txt:
            matched_standards.append("GB 50187")
        if "50058" in txt:
            matched_standards.append("GB 50058")
        if "50257" in txt:
            matched_standards.append("GB 50257")
        if "50444" in txt:
            matched_standards.append("GB 50444")
        if "50140" in txt:
            matched_standards.append("GB 50140")
        if "50016" in txt:
            matched_standards.append("GB 50016")
        if "45067" in txt:
            matched_standards.append("GB 45067")
        if "TSG 21" in txt:
            matched_standards.append("TSG 21")
        if "39800" in txt:
            matched_standards.append("GB 39800")
        if "JGJ 91" in txt:
            matched_standards.append("JGJ 91")
        if "HJ 2025" in txt:
            matched_standards.append("HJ 2025")
        
        if not matched_standards:
            matched_standards = ["OTHER / NO_STANDARD_KEYWORD"]
            
        for s in matched_standards:
            clusters[s].append(rec["hazardId"])

    print("\nCluster counts:")
    for s, hids in sorted(clusters.items(), key=lambda x: len(x[1]), reverse=True):
        print(f"  {s}: {len(hids)} hazards")

if __name__ == "__main__":
    main()

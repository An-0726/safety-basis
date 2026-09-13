# -*- coding: utf-8 -*-
"""把Excel交换稿拆成可审阅的现有实体补丁和新增实体候选，不直接改knowledge。"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
BOOK = Path(r"D:\Desktop\隐患库最终修订交付版_20260913.xlsx")
OUT = ROOT / "source" / "proposals" / "excel-20260913"
HAZARD_DIR = ROOT / "knowledge" / "hazards"

FIELDS = {
    "主题": "category",
    "隐患名称（标题）": "title",
    "隐患专业描述": "description",
    "整改措施": "measures",
    "适用条件": "conditions",
    "备注": "note",
}


def text(value):
    return "" if value is None else str(value).replace("\r\n", "\n").strip()


hazards = {}
for path in HAZARD_DIR.glob("*.json"):
    row = json.loads(path.read_text(encoding="utf-8"))
    hazards[row["id"]] = row

wb = load_workbook(BOOK, read_only=True, data_only=True)
ws = wb["建议回灌"]
headers = [text(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
idx = {name: i for i, name in enumerate(headers)}

existing = []
new = []
for values in ws.iter_rows(min_row=2, values_only=True):
    hid = text(values[idx["隐患ID"]])
    row = {
        "hazardId": hid,
        "sourceRow": int(values[idx["序号"]] or 0),
        "basisReview": text(values[idx["依据审查结论"]]),
        "finalDecision": text(values[idx["最终处理结论"]]),
        "publishable": text(values[idx["可发布"]]),
        "incomingPlaceScope": text(values[idx["场所"]]),
        "directBasis": text(values[idx["直接依据"]]),
        "basisQuote": text(values[idx["依据原文"]]),
        "fallbackBasis": text(values[idx["补充/兜底依据"]]),
        "changedFields": {},
    }
    for excel_name, json_name in FIELDS.items():
        row["changedFields"][json_name] = text(values[idx[excel_name]])
    if hid in hazards:
        old = hazards[hid]
        row["existingSnapshot"] = {name: old.get(name) for name in FIELDS.values()}
        row["existingPlaces"] = old.get("places") or []
        row["incomingPlaceScopeNeedsDecision"] = True
        existing.append(row)
    else:
        row["incomingPlaceScopeNeedsDecision"] = True
        new.append(row)

OUT.mkdir(parents=True, exist_ok=True)
(OUT / "existing-hazard-updates.json").write_text(json.dumps({
    "kind": "excel-existing-hazard-update-proposal-v1",
    "source": str(BOOK),
    "count": len(existing),
    "placeFieldPolicy": "not-applied; Excel场所为宽泛适用范围，knowledge.places为检索标签，需人工确定映射",
    "rows": existing,
}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
(OUT / "new-hazard-candidates.json").write_text(json.dumps({
    "kind": "excel-new-hazard-candidate-v1",
    "source": str(BOOK),
    "count": len(new),
    "status": "candidate-only; not admitted to knowledge",
    "rows": new,
}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
(OUT / "summary.json").write_text(json.dumps({
    "existingHazardUpdates": len(existing),
    "newHazardCandidates": len(new),
    "placeMappingPending": len(existing) + len(new),
    "source": str(BOOK),
}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"existingHazardUpdates": len(existing), "newHazardCandidates": len(new), "output": str(OUT)}, ensure_ascii=False, indent=2))

# -*- coding: utf-8 -*-
"""读取权威 Excel 并分析 434 条 proposed 的原始字段数据。"""
from __future__ import annotations
import json
import sys
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[2]
KNOW = ROOT / "knowledge"
EXCEL_PATH = Path(r"D:\Desktop\隐患库_1929条_新版口径全部整改完成_20260914.xlsx")

def main():
    proposed_ids = set(
        p.stem for p in (KNOW / "hazards").glob("*.json")
        if json.loads(p.read_text(encoding="utf-8")).get("lifecycle") == "proposed"
    )
    print(f"Total proposed hazards in knowledge: {len(proposed_ids)}")

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sheet = wb["隐患明细_修订后"]
    headers = [cell.value for cell in sheet[1]]
    print("Excel Headers:", headers)

    id_col_idx = headers.index("隐患ID") if "隐患ID" in headers else 0
    print(f"ID column index: {id_col_idx}")

    excel_rows = {}
    for r in range(2, sheet.max_row + 1):
        row_vals = [sheet.cell(r, c).value for c in range(1, len(headers) + 1)]
        hid = str(row_vals[id_col_idx]).strip() if row_vals[id_col_idx] else ""
        if hid:
            excel_rows[hid] = dict(zip(headers, row_vals))

    print(f"Total valid rows in Excel: {len(excel_rows)}")

    matched = sum(1 for hid in proposed_ids if hid in excel_rows)
    print(f"Proposed matched in Excel: {matched} / {len(proposed_ids)}")

    unmatched = [hid for hid in proposed_ids if hid not in excel_rows]
    print(f"Proposed not in Excel (should be extra non-target): {unmatched}")

    # 抽取 434 条 proposed 在 Excel 中的“修订直接依据”与“修订说明”
    proposed_in_excel = {hid: excel_rows[hid] for hid in proposed_ids if hid in excel_rows}
    
    # 打印前 5 条样本
    for i, (hid, row) in enumerate(list(proposed_in_excel.items())[:5]):
        print(f"\n--- Sample {i+1}: {hid} ---")
        for k in ["隐患名称", "类别", "修订状态", "修订直接依据", "修订说明", "依据条款要点"]:
            print(f"  {k}: {row.get(k)}")

if __name__ == "__main__":
    main()

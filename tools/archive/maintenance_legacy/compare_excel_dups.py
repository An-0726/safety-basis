# -*- coding: utf-8 -*-
import openpyxl
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
EXCEL_PATH = Path(r"D:\Desktop\隐患库_1929条_新版口径全部整改完成_20260914.xlsx")
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
sheet = wb["隐患明细_修订后"]
headers = [cell.value for cell in sheet[1]]
id_idx = headers.index("隐患ID")
for r in range(2, sheet.max_row + 1):
    val = sheet.cell(r, id_idx + 1).value
    if str(val).strip() in ["H_5E36D9ABBA9321D566592E49_2", "H_66B2A0967E8B4E4BAD7749DF_1"]:
        row_vals = [sheet.cell(r, c).value for c in range(1, len(headers) + 1)]
        d = dict(zip(headers, row_vals))
        print("--------------------------------------------------")
        print("ID:", d["隐患ID"])
        print("直接依据:", d["直接依据"])
        print("条款要点:", d["依据条款要点（修订，非原文摘录）"])
        print("修订说明:", d["修订说明"])
        print("整改措施:", d["整改措施"])

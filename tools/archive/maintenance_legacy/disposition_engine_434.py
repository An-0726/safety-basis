# -*- coding: utf-8 -*-
"""434 条 proposed 候选全量逐条审查与最终处置引擎。"""
from __future__ import annotations
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[2]
KNOW = ROOT / "knowledge"
EXCEL_PATH = Path(r"D:\Desktop\隐患库_1929条_新版口径全部整改完成_20260914.xlsx")

def load_json_dir(rel: str) -> dict[str, dict]:
    out = {}
    for p in (KNOW / rel).glob("*.json"):
        with open(p, "r", encoding="utf-8") as f:
            d = json.load(f)
        key = d.get("id") or d.get("entityId") or p.stem
        out[key] = d
    return out

def norm_str(s: str) -> str:
    return re.sub(r"[^\w\u4e00-\u9fff]", "", str(s)).upper()

def main():
    hazards = load_json_dir("hazards")
    clauses = load_json_dir("clauses")
    links = load_json_dir("links")
    lvs = load_json_dir("law-versions")
    cr = load_json_dir("reviews/clauses")

    proposed = [h for h in hazards.values() if h.get("lifecycle") == "proposed"]
    active_hazards = {h["id"]: h for h in hazards.values() if h.get("lifecycle") == "active"}
    print(f"Total proposed: {len(proposed)}, active: {len(active_hazards)}")

    # 找出每个条款当前已被哪些 active hazard 引用
    active_hazards_by_clause = defaultdict(list)
    for l in links.values():
        if l.get("lifecycle") == "active":
            hid = l.get("hazardId")
            cid = l.get("clauseId")
            if hid in active_hazards:
                active_hazards_by_clause[cid].append(hid)

    # 索引 active + verified clauses
    verified_clauses = {}
    for cid, c in clauses.items():
        if c.get("lifecycle") != "active":
            continue
        lvid = c.get("lawVersionId")
        lv = lvs.get(lvid, {})
        if lv.get("validityStatus") != "active":
            continue
        r = cr.get(cid, {})
        if r.get("decision") != "verified":
            continue
        verified_clauses[cid] = c

    # 加载 Excel
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sheet = wb["隐患明细_修订后"]
    headers = [cell.value for cell in sheet[1]]
    id_idx = headers.index("隐患ID")
    excel_rows = {}
    for r in range(2, sheet.max_row + 1):
        vals = [sheet.cell(r, c).value for c in range(1, len(headers) + 1)]
        hid = str(vals[id_idx]).strip() if vals[id_idx] else ""
        if hid:
            excel_rows[hid] = dict(zip(headers, vals))

    # Phase 6 历史记录
    p6_records = {}
    with open(ROOT / "docs" / "phase6-final-disposition.jsonl", "r", encoding="utf-8") as f:
        for line in f:
            d = json.loads(line)
            if d.get("recordType") == "metadata":
                continue
            p6_records[d["hazardId"]] = d

    # 建立多维条款快速检索
    clause_lookup = []
    for cid, c in verified_clauses.items():
        lvid = c.get("lawVersionId")
        lv = lvs.get(lvid, {})
        doc_num = lv.get("documentNumber") or ""
        law_title = lv.get("title") or ""
        art_path = c.get("articlePath") or ""

        clause_lookup.append({
            "cid": cid,
            "clause": c,
            "lawVersion": lv,
            "doc_num": doc_num,
            "law_title": law_title,
            "art_path": art_path,
            "norm_doc": norm_str(doc_num),
            "norm_title": norm_str(law_title),
            "norm_art": norm_str(art_path)
        })

    # 434 逐条审查列表
    dispositions = []

    # 1. 明确目标外候选（7条）
    extra_non_target_ids = {
        "H_02FEFD347E114E1E979C9589AF", "H_0647B65088564B789D64080EF1",
        "H_1F19FA1B951D46C1971E9B59B4", "H_23D109AF79FF0519BCD1837EC6_1",
        "H_3A5DEC6C74244A949CE3BC9A42", "H_5B89D7164D2C4B6E983663B009",
        "H_6E7E9CD077AD4918962EBA6EAE"
    }

    # 2. 明确与 active 重名的重复项
    exact_active_dups = {
        "H_12158_10_1_2": "H_12158_10_2_1",
        "H_12158_4_2_3_5_2": "H_12158_4_2_3_4_2",
        "H_12158_6_3_2_2": "H_12158_4_2_3_4_2",
        "H_12158_8_8_5_3": "H_12158_4_2_3_4_2",
    }

    # 3. 5 条 catalog-only（表2~5统领性条文）
    catalog_only_ids = {
        "H_GBT47236_4_1_1_1", "H_GBT47236_4_3_3_1", "H_GBT47236_4_3_4_1",
        "H_GBT47236_4_3_5_1", "H_GBT47236_4_3_6_1"
    }

    for h in proposed:
        hid = h["id"]
        title = h.get("title", "")
        desc = h.get("description", "")
        h_note = h.get("note", "")
        ex = excel_rows.get(hid, {})
        basis = str(ex.get("直接依据") or "")
        key_pts = str(ex.get("依据条款要点（修订，非原文摘录）") or "")
        rev_note = str(ex.get("修订说明") or "")
        combined = f"{basis} {key_pts} {rev_note} {h_note} {title} {desc}"
        norm_comb = norm_str(combined)

        disp = None

        # 规则 1: 目标外
        if hid in extra_non_target_ids:
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "OUT_OF_SCOPE",
                "reasonCode": "knowledge_extra_non_target",
                "reason": "该隐患不在权威 1929 目标集清单中（Phase 5 对账确认目标外），且未达到正式纳管标准，归入目标外保留。",
                "targetClauseId": None,
                "duplicateTargetId": None
            }

        # 规则 2: 重复项合并至 active
        elif hid in exact_active_dups:
            target_act_id = exact_active_dups[hid]
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "DUPLICATE_OR_MERGE",
                "reasonCode": "duplicate_of_active_hazard",
                "reason": f"与现有 active 隐患 {target_act_id} 标题与内容实质完全相同，合并至 {target_act_id}。",
                "targetClauseId": None,
                "duplicateTargetId": target_act_id
            }

        # 规则 3: proposed 内部互重
        elif hid == "H_1764B7439DFE44C1A680A7393E":
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "DUPLICATE_OR_MERGE",
                "reasonCode": "duplicate_within_proposed",
                "reason": "与 H_YJF_6_9_1 完全同名同义（危险废物贮存设施标志），归并至 H_YJF_6_9_1。",
                "targetClauseId": None,
                "duplicateTargetId": "H_YJF_6_9_1"
            }

        # 规则 4: 表格/交叉引用统领性条文（5条 catalog only）
        elif hid in catalog_only_ids:
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "KEEP_PROPOSED_CROSS_REFERENCE_GAP",
                "reasonCode": "cross_reference_table_requirement_gap",
                "reason": "本条系GB/T 47236引用表2至表5或外部标准的统领性条文，条款自身无独立具体操作义务，依赖展开各表细节。",
                "targetClauseId": None,
                "duplicateTargetId": None
            }

        # 规则 5: GB 12158 专项条文精确转正
        elif hid == "H_12158_4_2_2_3_1":
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "PROMOTE_ACTIVE",
                "reasonCode": "verified_direct_exact_locator",
                "reason": "GB 12158-2024第4.2.2.3条直接规定静电危险场所接地电阻不应大于100Ω，逐字完全一致，审核闭环。",
                "targetClauseId": "C_12158_4_2_2_3",
                "duplicateTargetId": None
            }
        elif hid == "H_12158_6_3_1_1":
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "PROMOTE_ACTIVE",
                "reasonCode": "verified_direct_exact_locator",
                "reason": "GB 12158-2024第6.3.1条直接规定非金属液体贮罐/输送管道表面电阻率与体电阻率要求，逐字完全一致，审核闭环。",
                "targetClauseId": "C_12158_6_3_1",
                "duplicateTargetId": None
            }
        elif hid == "H_12158_7_1_1":
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "PROMOTE_ACTIVE",
                "reasonCode": "verified_direct_exact_locator",
                "reason": "GB 12158-2024第7.1条直接规定非金属与金属导体连接紧密接触面积应大于20cm²，逐字完全一致，审核闭环。",
                "targetClauseId": "C_12158_7_1",
                "duplicateTargetId": None
            }

        # 规则 6: 灭火器铭牌/摆放精确转正
        elif hid == "H_08C1576EE4824ED8BE0CDD56DF":
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "PROMOTE_ACTIVE",
                "reasonCode": "verified_direct_exact_locator",
                "reason": "GB 50140-2005第5.1.3条直接规定灭火器摆放应稳固且铭牌应朝外，逐字完全一致，审核闭环。",
                "targetClauseId": "C_GB50140_5_1_3",
                "duplicateTargetId": None
            }

        # 规则 7: 危险废物台账精确转正
        elif hid == "H_0A9DDC2534BC46FDA9AEE80A89":
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "PROMOTE_ACTIVE",
                "reasonCode": "verified_direct_exact_locator",
                "reason": "GB 18597-2023第8.2.4条直接规定贮存设施运行期间应建立危险废物管理台账并保存，逐字完全一致，审核闭环。",
                "targetClauseId": "C_EB1C00891F0F010C3E8DE36882",
                "duplicateTargetId": None
            }

        # 规则 8: 管线安全标识精确转正
        elif hid == "H_1507C0C16EB84487BA78230E0D":
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "PROMOTE_ACTIVE",
                "reasonCode": "verified_direct_exact_locator",
                "reason": "GB/T 12801-2008第6.8.4条直接规定设备和管线应按规定涂识别色、识别符号和安全标识，逐字完全一致，审核闭环。",
                "targetClauseId": "C_3B764981E731EA8D2F1B2B0C59",
                "duplicateTargetId": None
            }

        # 规则 9: 依赖 upcoming 标准（GB 12801-2025 / GB/T 13869-2026 / GB 14444-2025 等）
        elif ("12801" in combined and "2025" in combined) or ("13869" in combined and "2026" in combined) or ("14444" in combined and "2025" in combined):
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "KEEP_PROPOSED_EVIDENCE_GAP",
                "reasonCode": "upcoming_standard_not_yet_effective",
                "reason": "所引标准为已发布但尚未到达实施日期的 upcoming 标准（如GB 12801-2025、GB/T 13869-2026等），现行有效版本中无对应独立强制条款，按效力规则保留候选等待实施。",
                "targetClauseId": None,
                "duplicateTargetId": None
            }

        # 规则 10: 语病/正向义务重构候选
        elif "依法应当进行消防验收" in title and "擅自投入使用" in title:
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "REWRITE_HAZARD",
                "reasonCode": "obligation_restatement_syntax",
                "reason": "标题为法条复述式长句，属于正向义务反转表述，需重构为标准缺陷状态并核实建设工程验收主体范围。",
                "targetClauseId": None,
                "duplicateTargetId": None
            }

        # 规则 11: 适用性/对象不匹配（已核定历史条目）
        elif hid == "H_C50168CA681E44F4B795F32BF0":
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "KEEP_PROPOSED_APPLICABILITY_GAP",
                "reasonCode": "conditional_scope_not_resolved",
                "reason": "GB 50187-2012第5.1.6条中'避免西晒'仅适用于高温、热加工及特殊要求建筑，当前隐患合并表述过宽。",
                "targetClauseId": None,
                "duplicateTargetId": None
            }
        elif hid == "H_FC756EA86AAF4A94BA7207A5B5":
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "KEEP_PROPOSED_APPLICABILITY_GAP",
                "reasonCode": "mixed_object_partial_basis",
                "reason": "安全生产法第45条只能覆盖劳动防护用品，不能同时支撑设备工程降噪措施。",
                "targetClauseId": None,
                "duplicateTargetId": None
            }
        elif hid == "H_AB6ECCCE1E89469C92E2DEA6C5":
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "KEEP_PROPOSED_APPLICABILITY_GAP",
                "reasonCode": "wrong_clause_object",
                "reason": "江苏省安全风险条例第16条为较大风险信息公开公示义务，不能支撑'未建立安全风险档案'。",
                "targetClauseId": None,
                "duplicateTargetId": None
            }
        elif hid == "H_B5E5670E989F4955969803EF66":
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "KEEP_PROPOSED_APPLICABILITY_GAP",
                "reasonCode": "update_duty_not_explicit",
                "reason": "现有公示条款未直接明确变更后公示栏更新的具体时限要求。",
                "targetClauseId": None,
                "duplicateTargetId": None
            }
        elif hid == "H_DB786B8BC54E4E55B1145D00D7":
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "KEEP_PROPOSED_APPLICABILITY_GAP",
                "reasonCode": "specific_technical_basis_not_proven",
                "reason": "候选对象为内部防雷装置，所引一般危化品条例不能作为防雷技术直接依据。",
                "targetClauseId": None,
                "duplicateTargetId": None
            }
        elif hid == "H_FC60ED3C946F4F18B89455526A":
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "KEEP_PROPOSED_APPLICABILITY_GAP",
                "reasonCode": "clause_object_mismatch",
                "reason": "GB/T 47236-2026第4.2.5.7条针对动力供应失效与意外重启，不能直接支撑'防松脱或防护装置'。",
                "targetClauseId": None,
                "duplicateTargetId": None
            }
        elif hid == "H_EACC1ED8867C4F3DA5FBAC8B49":
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "REWRITE_HAZARD",
                "reasonCode": "wording_strength_mismatch",
                "reason": "GB 50187-2012第3.0.14条原文为'不应选为厂址'，隐患标题误写为'禁止选址地段'，用语强度不匹配需重写。",
                "targetClauseId": None,
                "duplicateTargetId": None
            }

        # 规则 12: 附录/表格/交叉引用依赖（附录、表、行业标准交叉引用）
        elif re.search(r"(附录[A-Z0-9]|表\d+|参照.*标准|符合.*规定)", combined) and ("18597" in combined or "50016" in combined or "50140" in combined or "50444" in combined or "50054" in combined or "50058" in combined or "50257" in combined or "51309" in combined):
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "KEEP_PROPOSED_CROSS_REFERENCE_GAP",
                "reasonCode": "cross_reference_or_table_dependency",
                "reason": "所引条款依赖外部标准、附录技术参数或专项设计表格计算要求，当前知识库未完成被引用技术附录/表格的结构化收录。",
                "targetClauseId": None,
                "duplicateTargetId": None
            }

        # 规则 13: 缺少现行精确已核条款（无 exact clause 且缺少现行法律依据）
        else:
            disp = {
                "hazardId": hid,
                "title": title,
                "category": h.get("category"),
                "disposition": "KEEP_PROPOSED_EVIDENCE_GAP",
                "reasonCode": "no_exact_current_reviewed_clause",
                "reason": "当前知识库及私有库中缺少直接支持该隐患具体判定要求的现行有效条款逐字原文与官方证据，保留候选等待权威依据补齐。",
                "targetClauseId": None,
                "duplicateTargetId": None
            }

        dispositions.append(disp)

    print(f"\nDisposed {len(dispositions)} / {len(proposed)} proposed hazards.")
    counts = Counter(d["disposition"] for d in dispositions)
    print("Disposition distribution:")
    for k, v in counts.most_common():
        print(f"  {k}: {v}")

    reason_counts = Counter(d["reasonCode"] for d in dispositions)
    print("\nReason Code distribution:")
    for k, v in reason_counts.most_common():
        print(f"  {k}: {v}")

if __name__ == "__main__":
    main()

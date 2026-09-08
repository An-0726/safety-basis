"""Route mixed source files to the existing intake parser by sheet and headers."""
import argparse
import csv
import json
from pathlib import Path
import shutil

import intake

ROOT = Path(__file__).resolve().parents[2]
EXTENSIONS = intake.EXTENSIONS


def _archive(path, archive):
    blob = path.read_bytes()
    sha = intake.digest(blob)
    dest = Path(archive) / sha / "original"
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists():
        if intake.digest(dest.read_bytes()) != sha:
            raise ValueError(f"archive checksum mismatch: {path}")
    else:
        shutil.copyfile(path, dest)
    return sha, str(dest.resolve())


def _headers(path, mappings):
    ext = path.suffix.lower()
    if ext == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            height = max([m.get("headerRow", 1) for m in mappings] + [1])
            return [(name, [list(row) for row in book[name].iter_rows(
                max_row=height, values_only=True)]) for name in book.sheetnames]
        finally:
            book.close()
    if ext == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        return [("CSV", rows)]
    if ext == ".json":
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        if isinstance(data, dict):
            for mapping in mappings:
                if mapping.get("recordsKey") in data:
                    data = data[mapping["recordsKey"]]
                    break
        return [("JSON", [list(data[0]) if data and isinstance(data[0], dict) else []])]
    raise ValueError(f"unsupported format: {path.suffix}")


def _norm_headers(row):
    return {intake.norm(x) for x in (row or []) if intake.norm(x)}


def _matches(layout, mapping, sheet, rows):
    allowed = layout.get("sheetNames", [])
    if allowed and sheet not in allowed:
        return False
    if sheet in layout.get("forbidSheetNames", []):
        return False
    row_no = mapping.get("headerRow", 1)
    if row_no < 1 or row_no > len(rows):
        return False
    headers = _norm_headers(rows[row_no - 1])
    combos = layout.get("headerSets", [])
    if combos and not any({intake.norm(x) for x in combo} <= headers for combo in combos):
        return False
    contains = {intake.norm(x) for x in layout.get("headerContains", [])}
    if contains and not all(any(token in header for header in headers) for token in contains):
        return False
    forbidden = {intake.norm(x) for x in layout.get("forbidHeaders", [])}
    return not (headers & forbidden)


def load_registry(path):
    registry = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    base = Path(path).parent
    layouts = []
    for layout in registry.get("layouts", []):
        item = dict(layout)
        mapping_path = base / item["mapping"]
        item["mappingPath"] = mapping_path
        item["mappingData"] = json.loads(mapping_path.read_text(encoding="utf-8-sig"))
        layouts.append(item)
    for skip in registry.get("auxiliarySheets", []):
        item = dict(skip)
        item["kind"] = "auxiliary"
        item["mappingData"] = {"headerRow": item.get("headerRow", 1)}
        layouts.append(item)
    if not layouts:
        raise ValueError("registry has no layouts")
    return layouts


def route(path, layouts):
    mappings = [x["mappingData"] for x in layouts]
    units = _headers(path, mappings)
    routes, diagnostics, skipped = [], [], []
    for sheet, rows in units:
        matches = [x for x in layouts if _matches(x, x["mappingData"], sheet, rows)]
        headers = sorted(_norm_headers(rows[0] if rows else []))
        data = [x for x in matches if x.get("kind", "data") == "data"]
        auxiliary = [x for x in matches if x.get("kind") == "auxiliary"]
        if len(data) == 1 and not auxiliary:
            routes.append((sheet, data[0]))
        elif not data and len(auxiliary) == 1:
            skipped.append({"sheet": sheet, "auxiliaryId": auxiliary[0]["id"], "headers": headers, "reason": auxiliary[0].get("skipReason", "registered auxiliary sheet")})
        elif not matches:
            diagnostics.append({"sheet": sheet, "code": "unknown-layout", "message": "no registered Sheet/header combination", "headers": headers})
        else:
            diagnostics.append({"sheet": sheet, "code": "ambiguous-layout", "message": "multiple registered layouts", "layoutIds": [x["id"] for x in matches], "headers": headers})
    return routes, diagnostics, skipped


def process_file(path, layouts, db, law_db, archive):
    path = Path(path).resolve()
    sha, archive_path = _archive(path, archive)
    try:
        routes, diagnostics, skipped = route(path, layouts)
    except Exception as exc:
        return {"file": str(path), "sha256": sha, "status": "rejected", "routes": [], "skippedSheets": [], "diagnostics": [{"code": "inspect-error", "message": str(exc)}], "archivePath": archive_path}
    if diagnostics:
        return {"file": str(path), "sha256": sha, "status": "rejected", "routes": [], "skippedSheets": skipped, "plannedRoutes": [x["id"] for _, x in routes], "diagnostics": diagnostics, "archivePath": archive_path}
    results = []
    for sheet, layout in routes:
        mapping = dict(layout["mappingData"])
        if path.suffix.lower() == ".xlsx":
            mapping["sheet"] = sheet
        target = law_db if layout.get("database") == "law-registers" else db
        connection = intake.connect(target)
        try:
            result = intake.ingest(connection, path, mapping, archive)
        finally:
            connection.close()
        results.append({"sheet": sheet, "layoutId": layout["id"], "database": str(target), "result": result})
    failed = any(x["result"]["status"] != "complete" for x in results)
    return {"file": str(path), "sha256": sha, "status": "rejected" if failed else "complete", "routes": results, "skippedSheets": skipped, "diagnostics": [], "archivePath": archive_path}


def process_files(inputs, registry, db, law_db, archive):
    layouts = load_registry(registry)
    files = set()
    for item in map(Path, inputs):
        if item.is_dir():
            files.update(x.resolve() for x in item.rglob("*") if x.is_file() and x.suffix.lower() in EXTENSIONS and not x.name.startswith("~$"))
        elif item.is_file():
            files.add(item.resolve())
        else:
            raise ValueError(f"input does not exist: {item}")
    return [process_file(path, layouts, db, law_db, archive) for path in sorted(files)]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("inputs", nargs="+")
    parser.add_argument("--registry", type=Path, default=ROOT / "source/mappings/registry.json")
    parser.add_argument("--db", type=Path, default=ROOT / "source/staging/intake.sqlite3")
    parser.add_argument("--law-db", type=Path, default=ROOT / "source/staging/law-registers.sqlite3")
    parser.add_argument("--archive", type=Path, default=ROOT / "source/archive")
    args = parser.parse_args()
    results = process_files(args.inputs, args.registry, args.db, args.law_db, args.archive)
    print(json.dumps(results, ensure_ascii=False, indent=2))
    return int(any(x["status"] != "complete" for x in results))


if __name__ == "__main__":
    raise SystemExit(main())

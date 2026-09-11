"""Extract review candidates from numbered standards without publishing them.

This adapter is deliberately separate from ``legal_text.py``.  Standards use
decimal clause locators and often place requirements in tables; the output is
an auditable staging artifact, not a legal-verification decision and not a
catalog import request.
"""
from __future__ import annotations

import argparse
from io import BytesIO
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

import exchange
import legal_text
import review


LOCATOR = re.compile(r"^\s*([1-9]\d*(?:\.\d+)+)\s+(.+?)\s*$")
CHAPTER = re.compile(r"^\s*([1-9]\d*)\s+(.+?)\s*$")
ANNEX = re.compile(r"^\s*附\s*录\s+([A-Z])\s*$", re.I)
TABLE = re.compile(r"^\s*表\s*([A-Z]?\.?\d+)(?:\s+(.+?))?(?:\s*[（(]?续[）)]?)?\s*$", re.I)
DOC_HEADER = re.compile(r"^\s*(?:GB|GB/T|AQ|HG|DL|JB)(?:\s|/|\d).*[—-]\d{4}\s*$", re.I)
PAGE_NUMBER = re.compile(r"^\s*\d{1,3}\s*$")
REFERENCES = re.compile(r"^\s*参\s*考\s*文\s*献\s*$")


def _clean(text: str) -> str:
    return re.sub(r"[ \t\u3000]+", " ", text).strip()


def _content_rows(blob: bytes) -> list[dict[str, Any]]:
    raw = legal_text.paragraphs(blob)
    rows: list[dict[str, Any]] = []
    for ordinal, value in enumerate(raw):
        text = _clean(value)
        if not text or DOC_HEADER.match(text):
            continue
        # PDF extraction commonly reverses printed page 14 into ``41``.  A
        # standalone number immediately before a repeated document header is
        # page furniture, while numbers inside a table are retained.
        next_text = _clean(raw[ordinal + 1]) if ordinal + 1 < len(raw) else ""
        if PAGE_NUMBER.match(text) and DOC_HEADER.match(next_text):
            continue
        rows.append({"paragraph": ordinal, "text": text})
    return rows


def _section_heading(text: str) -> tuple[str, str] | None:
    match = LOCATOR.match(text)
    if not match:
        return None
    locator, tail = match.groups()
    # Table cells frequently begin with a bare cross-reference.  A heading or
    # requirement starts with meaningful text rather than punctuation.
    if not re.search(r"[\u4e00-\u9fffA-Za-z]", tail) or tail[0] in "、，,；;。)）":
        return None
    return locator, tail


def _looks_like_toc(text: str) -> bool:
    return "…" in text or bool(re.search(r"\.{5,}", text))


def _next_starts_subsection(rows: list[dict[str, Any]], position: int, chapter: str) -> bool:
    for row in rows[position + 1:position + 5]:
        heading = _section_heading(row["text"])
        if heading and heading[0].startswith(chapter + "."):
            return True
    return False


def _table_follows(rows: list[dict[str, Any]], position: int) -> bool:
    return any(TABLE.match(row["text"]) for row in rows[position + 1:position + 8])


def directory(blob: bytes, *, chapters: set[str] | None = None,
              annexes: set[str] | None = None) -> dict[str, Any]:
    chapters = chapters or {"4", "5", "6"}
    annexes = {value.upper() for value in (annexes or {"B"})}
    rows = _content_rows(blob)
    nodes: list[dict[str, Any]] = []
    active_scope: str | None = None
    active_annex: str | None = None
    current: dict[str, Any] | None = None
    table: dict[str, Any] | None = None

    def finish_current() -> None:
        nonlocal current
        if current is None:
            return
        quote = "\n".join(current.pop("_parts"))
        current["quote"] = quote
        current["textSha256"] = exchange.sha256_bytes(quote.encode("utf-8"))
        nodes.append(current)
        current = None

    def finish_table() -> None:
        nonlocal table
        if table is None:
            return
        text = "\n".join(table.pop("_parts"))
        table["rawText"] = text
        table["textSha256"] = exchange.sha256_bytes(text.encode("utf-8"))
        table["rowParsingStatus"] = "待人工确认"
        nodes.append(table)
        table = None

    for position, row in enumerate(rows):
        ordinal, text = row["paragraph"], row["text"]
        if _looks_like_toc(text):
            continue
        if REFERENCES.match(text):
            finish_current()
            finish_table()
            active_scope = None
            active_annex = None
            continue
        annex = ANNEX.match(text)
        if annex:
            finish_current()
            finish_table()
            active_annex = annex.group(1).upper()
            active_scope = f"附录 {active_annex}" if active_annex in annexes else None
            if active_scope:
                nodes.append({"nodeType": "annex", "locator": active_scope,
                              "title": "", "paragraphStart": ordinal,
                              "paragraphEnd": ordinal})
            continue

        chapter = CHAPTER.match(text)
        if chapter and "." not in chapter.group(1):
            number, title = chapter.groups()
            if table is not None:
                if number not in chapters or not _next_starts_subsection(rows, position, number):
                    table["paragraphEnd"] = ordinal
                    table["_parts"].append(text)
                    continue
            if number in chapters:
                finish_current()
                finish_table()
                active_annex = None
                active_scope = number
                nodes.append({"nodeType": "chapter", "locator": number,
                              "title": title, "paragraphStart": ordinal,
                              "paragraphEnd": ordinal})
                continue
            if active_annex is None:
                finish_current()
                finish_table()
                active_scope = None

        if active_scope is None:
            continue

        heading = _section_heading(text)
        # A new subsection at depth <= 3 closes the preceding standards table.
        # Deeper decimal values inside a table are usually cross-references.
        if (table is not None and heading and len(heading[0].split(".")) <= 3
                and _table_follows(rows, position)):
            finish_table()

        table_match = TABLE.match(text)
        if table_match:
            finish_current()
            label = table_match.group(1).upper()
            title = table_match.group(2) or ""
            if table is None or table["locator"] != f"表 {label}":
                finish_table()
                table = {"nodeType": "table", "locator": f"表 {label}",
                         "title": title, "scope": active_scope,
                         "paragraphStart": ordinal, "paragraphEnd": ordinal,
                         "_parts": [text]}
            else:
                table["paragraphEnd"] = ordinal
                table["_parts"].append(text)
            continue

        if table is not None:
            table["paragraphEnd"] = ordinal
            table["_parts"].append(text)
            continue

        if heading:
            locator, first_text = heading
            if locator.split(".")[0] not in chapters:
                continue
            finish_current()
            current = {"nodeType": "clause_candidate", "locator": locator,
                       "scope": active_scope, "paragraphStart": ordinal,
                       "paragraphEnd": ordinal, "_parts": [text],
                       "firstLineText": first_text,
                       "verificationStatus": "待核验"}
        elif current is not None:
            current["paragraphEnd"] = ordinal
            current["_parts"].append(text)

    finish_current()
    finish_table()
    locators = [node["locator"] for node in nodes if node["nodeType"] == "clause_candidate"]
    duplicates = sorted({value for value in locators if locators.count(value) > 1})
    return {
        "formatVersion": "safety-standard-directory-v1",
        "snapshotSha256": exchange.sha256_bytes(blob),
        "sourceParagraphCount": len(legal_text.paragraphs(blob)),
        "selectedChapters": sorted(chapters),
        "selectedAnnexes": sorted(annexes),
        "candidateCount": len(locators),
        "tableCount": sum(node["nodeType"] == "table" for node in nodes),
        "extractionIssues": (["Duplicate clause locators: " + ", ".join(duplicates)] if duplicates else []),
        "legalVerification": "待核验",
        "catalogImportStatus": "未生成",
        "nodes": nodes,
    }


def write_review_workbook(data: dict[str, Any], output: Path, *, law_version_id: str = "") -> None:
    """Write a human/agent review surface without making it an import format."""
    output = Path(output).resolve()
    exchange.check_output(output, ".xlsx")
    wb = Workbook()
    guide = wb.active
    guide.title = "说明"
    guide.append(["标准条款抽取审阅表"])
    guide.append(["法规版本ID", law_version_id])
    guide.append(["原件SHA-256", data["snapshotSha256"]])
    guide.append(["候选条款数", data["candidateCount"]])
    guide.append(["表格块数", data["tableCount"]])
    guide.append(["使用规则", "本文件仅供审阅；不得直接导入母库。确认后的记录另行生成 catalog-v1 请求并再次通过 review.py 门禁。"])
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 100
    guide.sheet_view.showGridLines = False

    clauses = wb.create_sheet("条款候选")
    clause_headers = ["法规版本ID", "条款定位", "原文候选", "原段落起", "原段落止",
                      "文本SHA-256", "审阅结论", "审阅备注"]
    clauses.append(clause_headers)
    for node in data["nodes"]:
        if node["nodeType"] != "clause_candidate":
            continue
        values = [law_version_id, node["locator"], node["quote"], node["paragraphStart"],
                  node["paragraphEnd"], node["textSha256"], "待核验", ""]
        for column, value in enumerate(values, 1):
            exchange.write_cell(clauses, clauses.max_row + (1 if column == 1 else 0), column, value)
    decision = DataValidation(type="list", formula1='"待核验,通过,驳回,需修正"', allow_blank=False)
    clauses.add_data_validation(decision)
    decision.add(f"G2:G{clauses.max_row}")
    exchange._style_table(clauses, {7, 8})
    clauses.freeze_panes = "A2"
    clauses.column_dimensions["B"].width = 18
    clauses.column_dimensions["C"].width = 90
    clauses.column_dimensions["H"].width = 45

    tables = wb.create_sheet("表格块")
    table_headers = ["法规版本ID", "表格定位", "所属范围", "表格原文块", "原段落起", "原段落止",
                     "文本SHA-256", "行列解析状态", "审阅备注"]
    tables.append(table_headers)
    for node in data["nodes"]:
        if node["nodeType"] != "table":
            continue
        tables.append([law_version_id, node["locator"], node["scope"], node["rawText"],
                       node["paragraphStart"], node["paragraphEnd"], node["textSha256"],
                       node["rowParsingStatus"], ""])
    exchange._style_table(tables, {8, 9})
    tables.freeze_panes = "A2"
    tables.column_dimensions["D"].width = 100
    tables.column_dimensions["I"].width = 45

    buffer = BytesIO()
    wb.save(buffer)
    exchange.install_output(output, buffer.getvalue())


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--snapshot", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--chapters", nargs="+", default=["4", "5", "6"])
    parser.add_argument("--annexes", nargs="+", default=["B"])
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--law-version-id", default="")
    args = parser.parse_args()
    result = directory(args.snapshot.read_bytes(), chapters=set(args.chapters), annexes=set(args.annexes))
    review.write_json(args.output, result)
    if args.workbook:
        write_review_workbook(result, args.workbook, law_version_id=args.law_version_id)
    print({key: result[key] for key in ("candidateCount", "tableCount", "extractionIssues",
                                        "legalVerification", "catalogImportStatus")})
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

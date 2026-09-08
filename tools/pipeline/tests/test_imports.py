import json
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

PIPELINE = Path(__file__).resolve().parents[1]
REPO = PIPELINE.parents[1]
sys.path.insert(0, str(PIPELINE))
import imports  # noqa: E402


class ImportRoutingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.inputs = self.root / "imports"
        self.inputs.mkdir()
        self.archive = self.root / "archive"
        self.db = self.root / "staging" / "intake.sqlite3"
        self.law_db = self.root / "staging" / "law-registers.sqlite3"
        self.registry = REPO / "source" / "mappings" / "registry.json"

    def tearDown(self):
        self.temp.cleanup()

    def workbook(self, path):
        book = Workbook()
        first = book.active
        first.title = "普通检查"
        first.append(["隐患名称", "场所"])
        first.append(["灭火器失压", "车间"])
        task = book.create_sheet("隐患库任务")
        task.append(["隐患描述", "依据法规", "检查结果"])
        task.append(["通道被占用", "第三十五条", "不符合"])
        law = book.create_sheet("依据总表")
        law.append(["依据名称", "文号/标准号"])
        law.append(["安全生产法", "主席令第八十八号"])
        stats = book.create_sheet("分类统计")
        stats.append(["统计维度", "分类", "数量", "说明"])
        notes = book.create_sheet("更新与核验说明")
        notes.append(["某项目评价依据台账——核验说明", ""])
        diff = book.create_sheet("逐条核验差异")
        diff.append(["总表序号", "原依据名称", "核验后依据名称"])
        book.save(path)

    def count(self, db, table):
        conn = sqlite3.connect(db)
        try:
            return conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
        finally:
            conn.close()

    def test_mixed_sheets_route_and_repeat_without_new_rows(self):
        path = self.inputs / "mixed.xlsx"
        self.workbook(path)
        first = imports.process_files([path], self.registry, self.db, self.law_db, self.archive)
        self.assertEqual(first[0]["status"], "complete")
        self.assertEqual({x["layoutId"] for x in first[0]["routes"]}, {"default", "zcode-task", "law-register"})
        self.assertEqual({x["auxiliaryId"] for x in first[0]["skippedSheets"]}, {"law-category-statistics", "law-update-notes", "law-verification-diff"})
        self.assertEqual(self.count(self.db, "candidates"), 2)
        self.assertEqual(self.count(self.law_db, "candidates"), 1)
        second = imports.process_files([path], self.registry, self.db, self.law_db, self.archive)
        self.assertTrue(all(x["result"]["repeated"] for x in second[0]["routes"]))
        self.assertEqual(self.count(self.db, "candidates"), 2)
        self.assertEqual(self.count(self.law_db, "candidates"), 1)
        self.assertTrue(Path(first[0]["archivePath"]).exists())

    def test_equal_layout_candidates_are_rejected_with_diagnostic(self):
        path = self.inputs / "ambiguous.csv"
        path.write_text("隐患名称\n测试\n", encoding="utf-8")
        custom = self.root / "registry.json"
        mapping = REPO / "source" / "mappings" / "default.json"
        custom.write_text(json.dumps({"version": 1, "layouts": [
            {"id": "a", "mapping": str(mapping), "headerSets": [["隐患名称"]]},
            {"id": "b", "mapping": str(mapping), "headerSets": [["隐患名称"]]},
        ]}, ensure_ascii=False), encoding="utf-8")
        result = imports.process_files([path], custom, self.db, self.law_db, self.archive)[0]
        self.assertEqual(result["status"], "rejected")
        self.assertEqual(result["diagnostics"][0]["code"], "ambiguous-layout")
        self.assertTrue(Path(result["archivePath"]).exists())

    def test_unknown_headers_are_rejected_with_diagnostic(self):
        path = self.inputs / "unknown.csv"
        path.write_text("未知列\n值\n", encoding="utf-8")
        result = imports.process_files([path], self.registry, self.db, self.law_db, self.archive)[0]
        self.assertEqual(result["status"], "rejected")
        self.assertEqual(result["diagnostics"][0]["code"], "unknown-layout")
        self.assertTrue(Path(result["archivePath"]).exists())

    def test_unknown_business_sheet_is_not_silently_skipped(self):
        path = self.inputs / "unknown-business.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "依据总表"
        sheet.append(["依据名称", "文号/标准号"])
        sheet.append(["安全生产法", "主席令"])
        unknown = book.create_sheet("业务数据")
        unknown.append(["业务编号", "内部状态"])
        unknown.append(["A-1", "待处理"])
        book.save(path)
        result = imports.process_files([path], self.registry, self.db, self.law_db, self.archive)[0]
        self.assertEqual(result["status"], "rejected")
        self.assertEqual(result["diagnostics"][0]["sheet"], "业务数据")
        self.assertEqual(result["plannedRoutes"], ["law-register"])


if __name__ == "__main__":
    unittest.main()

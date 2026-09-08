"""Idempotent private intake only. No writes to content/, data/, or the master.

Raw files/rows are evidence, never instructions. Formula cells are preserved as
strings and flagged by the reader; formulas/macros are never evaluated.
"""
import argparse
import csv
import hashlib
import io
import json
from pathlib import Path
import sqlite3
import unicodedata
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

ROOT = Path(__file__).resolve().parents[2]
PARSER_VERSION = 'intake-v1'
EXTENSIONS = {'.csv', '.json', '.xlsx'}


def dumps(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'), default=str)


def digest(value):
    return hashlib.sha256(value if isinstance(value, bytes) else value.encode('utf-8')).hexdigest()


def norm(value):
    # Do not strip punctuation, negate words, units, article numbers or decimals.
    return ' '.join(unicodedata.normalize('NFKC', '' if value is None else str(value)).split())


def connect(db):
    db = Path(db)
    db.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db)
    conn.executescript((ROOT / 'source/schemas/intake.sql').read_text(encoding='utf-8'))
    return conn


def mapped(headers, values, mapping):
    if len(values) > len(headers) and any(x not in (None, '') for x in values[len(headers):]):
        raise ValueError('row has values beyond the header')
    named = [norm(x) for x in headers]
    if len([x for x in named if x]) != len(set(x for x in named if x)):
        raise ValueError('duplicate header: supply an unambiguous mapping/header row')
    result = {}
    for field, aliases in mapping['fields'].items():
        matches = [i for i, name in enumerate(named) if name and name in {norm(a) for a in aliases}]
        if len(matches) > 1:
            raise ValueError(f'ambiguous field {field}; select one source column in mapping')
        result[field] = norm(values[matches[0]]) if matches and matches[0] < len(values) else ''
    missing = [k for k in mapping.get('required', ['title']) if not result.get(k)]
    if missing:
        raise ValueError(f'missing fields: {missing}')
    return result


def table_rows(rows, mapping, sheet):
    header_no = mapping.get('headerRow', 1)
    if not isinstance(header_no, int) or header_no < 1:
        raise ValueError('headerRow must be a positive integer')
    headers = None
    for number, values in enumerate(rows, 1):
        values = list(values)
        if number == header_no:
            headers = values
        elif number > header_no and any(v not in (None, '') for v in values):
            raw = {'headers': headers, 'values': values}
            try:
                if any(isinstance(v, str) and v.startswith('=') for v in values):
                    raise ValueError('formula cell: supply reviewed static values; original retained')
                cleaned = mapped(headers, values, mapping)
            except ValueError as exc:
                yield sheet, number, raw, None, str(exc)
            else:
                yield sheet, number, raw, cleaned, ''
    if headers is None:
        raise ValueError(f'{sheet}: header row not found')


def read_rows(blob, extension, mapping):
    if extension == '.csv':
        rows = csv.reader(io.StringIO(blob.decode(mapping.get('encoding', 'utf-8-sig'))),
                          delimiter=mapping.get('delimiter', ','))
        yield from table_rows(rows, mapping, 'CSV')
    elif extension == '.json':
        data = json.loads(blob.decode('utf-8-sig'))
        if mapping.get('recordsKey'):
            data = data[mapping['recordsKey']]
        if not isinstance(data, list) or any(not isinstance(x, dict) for x in data):
            raise ValueError('JSON must contain an object array; set recordsKey for a wrapped array')
        for number, raw in enumerate(data, 1):
            try:
                cleaned = mapped(list(raw), list(raw.values()), mapping)
            except ValueError as exc:
                yield 'JSON', number, raw, None, str(exc)
            else:
                yield 'JSON', number, raw, cleaned, ''
    elif extension == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(blob), read_only=True, data_only=False, keep_links=False)
        try:
            names = [mapping['sheet']] if mapping.get('sheet') else book.sheetnames
            for name in names:
                yield from table_rows(book[name].iter_rows(values_only=True), mapping, name)
        finally:
            book.close()
    else:
        raise ValueError(f'unsupported format: {extension}')


def ingest(conn, file, mapping, archive):
    file, archive = Path(file), Path(archive)
    extension = file.suffix.lower()
    if extension not in EXTENSIONS:
        raise ValueError(f'unsupported format: {extension}')
    blob = file.read_bytes()  # Parse exactly these archived bytes, not a second read.
    sha = digest(blob)
    source_id = 'S_' + sha
    run_id = 'I_' + digest(dumps([sha, extension, PARSER_VERSION, mapping]))
    archived = archive / sha / 'original'
    archived.parent.mkdir(parents=True, exist_ok=True)
    if archived.exists():
        if digest(archived.read_bytes()) != sha:
            raise ValueError('archive checksum mismatch')
    else:
        with archived.open('xb') as stream:
            stream.write(blob)
    with conn:
        conn.execute('INSERT OR IGNORE INTO sources(id,sha256,archive_path,byte_size) VALUES(?,?,?,?)',
                     (source_id, sha, str(archived.resolve()), len(blob)))
        conn.execute('INSERT OR IGNORE INTO source_locations VALUES(?,?)', (source_id, str(file.resolve())))
    previous = conn.execute('SELECT status,row_count,error FROM import_runs WHERE id=?', (run_id,)).fetchone()
    if previous:
        return {'file': file.name, 'status': previous[0], 'rows': previous[1], 'repeated': True, 'error': previous[2]}
    rows, errors = [], []
    try:
        for sheet, number, raw, cleaned, error in read_rows(blob, extension, mapping):
            rows.append((sheet, number, raw, cleaned))
            if error:
                errors.append(f'{sheet}:{number}: {error}')
        if not rows:
            errors.append('no non-empty records')
    except (ValueError, KeyError, UnicodeError, csv.Error, BadZipFile, ParseError) as exc:
        errors.append(str(exc))
    status = 'rejected' if errors else 'complete'
    # One transaction for the complete parse. Rejected files retain raw rows,
    # but publish zero candidates; fixing the mapping creates a new run.
    with conn:
        conn.execute('INSERT INTO import_runs VALUES(?,?,?,?,?,?,?)',
                     (run_id, source_id, PARSER_VERSION, dumps(mapping), status, len(rows), dumps(errors)))
        for sheet, number, raw, cleaned in rows:
            row_id = 'R_' + digest(dumps([sha, sheet, number]))
            # Keep the first parse and each later mapping/encoding interpretation.
            conn.execute('INSERT OR IGNORE INTO source_rows VALUES(?,?,?,?,?)',
                         (row_id, source_id, sheet, number, dumps(raw)))
            conn.execute('INSERT INTO parsed_rows VALUES(?,?,?)', (run_id, row_id, dumps(raw)))
            if errors:
                continue
            identity = {k: v for k, v in cleaned.items() if k != 'externalId'}
            fp = digest(dumps(identity))
            candidate_id = 'D_' + fp
            conn.execute('INSERT OR IGNORE INTO candidates(id,fingerprint,normalized_json) VALUES(?,?,?)',
                         (candidate_id, fp, dumps(identity)))
            conn.execute('INSERT OR IGNORE INTO candidate_sources VALUES(?,?,?)', (candidate_id, row_id, run_id))
    return {'file': file.name, 'status': status, 'rows': len(rows), 'repeated': False, 'errors': errors}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('inputs', nargs='+', type=Path)
    parser.add_argument('--mapping', type=Path, default=ROOT / 'source/mappings/default.json')
    parser.add_argument('--db', type=Path, default=ROOT / 'source/staging/intake.sqlite3')
    parser.add_argument('--archive', type=Path, default=ROOT / 'source/archive')
    args = parser.parse_args()
    mapping = json.loads(args.mapping.read_text(encoding='utf-8-sig'))
    files = set()
    for item in args.inputs:
        if not item.exists():
            parser.error(f'input does not exist: {item}')
        if item.is_dir():
            files.update(x.resolve() for x in item.rglob('*')
                         if x.is_file() and x.suffix.lower() in EXTENSIONS and not x.name.startswith('~$'))
        else:
            files.add(item.resolve())
    if not files:
        parser.error('no supported files found')
    conn = connect(args.db)
    try:
        results = [ingest(conn, file, mapping, args.archive) for file in sorted(files)]
    finally:
        conn.close()
    print(json.dumps(results, ensure_ascii=False, indent=2))
    return int(any(x['status'] != 'complete' for x in results))


if __name__ == '__main__':
    raise SystemExit(main())

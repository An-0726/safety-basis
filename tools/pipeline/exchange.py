"""Versioned Excel exchange for the normalized safety master.

The workbook is an editable snapshot, never an independent source of truth.
Only existing core records can be updated in exchange-v1.  Importing first creates
a reviewable proposal; a separate command applies that exact proposal atomically.
"""
from __future__ import annotations

import argparse
from collections import Counter
import datetime as dt
import hashlib
from io import BytesIO
import json
import os
import sqlite3
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from master import ROOT, connect_readonly, dumps


FORMAT_VERSION = "safety-master-exchange-v1"

# Column definitions: (database column, workbook heading, editable, value kind).
SHEETS: dict[str, dict[str, Any]] = {
    "隐患库": {"table": "hazards", "entity": "hazard", "columns": [
        ("id", "ID", False, "text"), ("revision", "基础版本", False, "int"),
        ("operation", "操作", True, "operation"), ("title", "隐患标题", True, "text"),
        ("description", "隐患描述", True, "text"), ("measures", "整改措施", True, "text"),
        ("category", "分类", True, "text"), ("conditions", "适用条件", True, "text"),
        ("note", "备注", True, "text"), ("mode", "模式", True, "text"),
        ("status", "当前状态", False, "text"), ("checked", "原核验标记", False, "text"),
        ("merged_into", "合并至", False, "text"),
    ], "demote": {"status": "待核验", "checked": ""}, "required": ("title",)},
    "法规库": {"table": "laws", "entity": "law", "columns": [
        ("id", "ID", False, "text"), ("revision", "基础版本", False, "int"),
        ("operation", "操作", True, "operation"), ("canonical_name", "规范名称", True, "text"),
        ("issuer", "发布机关", True, "text"), ("jurisdiction_code", "适用地域", True, "text"),
        ("document_kind", "文种", True, "text"), ("identity_key", "身份键", False, "text"),
        ("identity_status", "身份状态", False, "text"), ("status", "当前状态", False, "text"),
        ("checked", "原核验标记", False, "text"),
    ], "demote": {"status": "待核验", "checked": ""}, "required": ("canonical_name",)},
    "法规版本": {"table": "law_versions", "entity": "law_version", "columns": [
        ("id", "ID", False, "text"), ("revision", "基础版本", False, "int"),
        ("operation", "操作", True, "operation"), ("law_id", "法规ID", False, "text"),
        ("version_key", "版本键", False, "text"), ("document_number", "发文字号", True, "text"),
        ("official_name", "正式名称", True, "text"), ("level", "效力层级", True, "text"),
        ("scope", "适用范围", True, "text"), ("effective_date", "生效日期", True, "date"),
        ("end_date", "失效日期", True, "date"), ("validity_status", "有效性判断", True, "validity"),
        ("source_url", "官方来源URL", True, "text"), ("review_status", "核验状态", False, "text"),
        ("checked", "原核验标记", False, "text"),
    ], "demote": {"review_status": "待核验", "checked": ""}, "required": ("official_name",)},
    "条款库": {"table": "clauses", "entity": "clause", "columns": [
        ("id", "ID", False, "text"), ("revision", "基础版本", False, "int"),
        ("operation", "操作", True, "operation"), ("law_version_id", "法规版本ID", False, "text"),
        ("article_path", "条款定位", True, "text"), ("quote", "条款原文", True, "text"),
        ("source_url", "官方来源URL", True, "text"), ("status", "当前状态", False, "text"),
        ("identity_status", "定位状态", False, "text"), ("checked", "原核验标记", False, "text"),
    ], "demote": {"status": "待核验", "checked": ""}, "required": ("article_path", "quote")},
    "依据关联": {"table": "links", "entity": "link", "columns": [
        ("id", "ID", False, "text"), ("revision", "基础版本", False, "int"),
        ("operation", "操作", True, "operation"), ("hazard_id", "隐患ID", False, "text"),
        ("clause_id", "条款ID", False, "text"), ("role", "依据角色", True, "text"),
        ("priority", "优先级", True, "int"), ("applicability", "适用性说明", True, "text"),
        ("jurisdiction_code", "适用地域", True, "text"), ("status", "当前状态", False, "text"),
    ], "demote": {"status": "待核验"}, "required": ()},
}

STATE_TABLES = (
    "sources", "source_locations", "batches", "source_rows", "hazards", "hazard_tags",
    "laws", "law_aliases", "law_versions", "law_successions", "clauses", "links",
    "legacy_payloads", "provenance", "migration_conflicts", "verification", "evidence",
)

READ_ONLY_SHEETS = {
    "隐患标签": ("hazard_tags", ("hazard_id", "kind", "value", "ordinal")),
    "法规别名": ("law_aliases", ("law_id", "alias", "ordinal")),
    "替代关系": ("law_successions", ("old_version_id", "new_version_id", "relation", "scope", "effective_date", "verification_id")),
    "原始来源": ("sources", ("id", "sha256", "original_name", "media_type", "byte_size", "visibility", "storage_ref", "source_commit", "created_at")),
    "核验记录": ("verification", ("id", "entity_type", "entity_id", "entity_revision", "check_type", "result", "reviewer", "model", "checked_at", "review_due_at", "evidence_id", "reason", "supersedes")),
    "来源关联": ("provenance", ("id", "entity_type", "entity_id", "source_row_id", "field_path", "transform_run_id")),
    "来源定位": ("source_rows", ("id", "source_id", "sheet", "row_number", "json_pointer")),
}


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def utc_now() -> str:
    return dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat()


def table_rows(conn: sqlite3.Connection, table: str) -> list[dict[str, Any]]:
    conn.row_factory = sqlite3.Row
    columns = [row[1] for row in conn.execute(f"PRAGMA table_info({table})")]
    order = ",".join(f'"{column}"' for column in columns)
    return [dict(row) for row in conn.execute(f'SELECT * FROM "{table}" ORDER BY {order}')]


def state_hash(conn: sqlite3.Connection) -> str:
    payload = {table: table_rows(conn, table) for table in STATE_TABLES}
    return sha256_bytes(dumps(payload).encode("utf-8"))


def normalize(value: Any, kind: str) -> Any:
    if value is None:
        if kind == "int":
            raise ValueError("整数列不能为空")
        return ""
    if kind == "int":
        if isinstance(value, bool):
            raise ValueError("布尔值不能作为整数")
        if isinstance(value, float) and not value.is_integer():
            raise ValueError(f"整数列出现小数: {value}")
        try:
            return int(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"无法解析整数: {value}") from exc
    if kind == "date" and isinstance(value, (dt.date, dt.datetime)):
        return value.date().isoformat() if isinstance(value, dt.datetime) else value.isoformat()
    if not isinstance(value, str):
        raise ValueError(f"文本列必须保留文本类型: {value}")
    return value


def check_output(output: Path, suffix: str) -> None:
    if output.suffix.lower() != suffix:
        raise ValueError(f"输出必须使用 {suffix} 扩展名")
    if output.exists():
        raise ValueError(f"输出已存在，请使用新文件名，避免覆盖编辑成果: {output}")


def install_output(output: Path, blob: bytes) -> None:
    """Install only complete artifacts, without replacing a concurrent writer."""
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix="exchange-", suffix=".tmp", dir=output.parent)
    try:
        with os.fdopen(fd, "wb") as stream:
            stream.write(blob)
        os.link(name, output)
    finally:
        Path(name).unlink(missing_ok=True)


def write_cell(ws, row: int, column: int, value: Any) -> None:
    if isinstance(value, str) and len(value) > 32767:
        raise ValueError(f"{ws.title}!{row}:{column}: 文本超过 Excel 单元格限制，拒绝截断")
    cell = ws.cell(row, column, value)
    if isinstance(value, str):
        # Literal text such as '=...' must never become executable Excel formulas.
        cell.data_type = "s"
        cell.number_format = "@"


def check_integrity(conn: sqlite3.Connection) -> None:
    if conn.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
        raise ValueError("母库完整性检查失败")
    if conn.execute("PRAGMA foreign_key_check").fetchall():
        raise ValueError("母库外键断链")


def _style_table(ws, editable_indices: set[int]) -> None:
    navy, amber, gray = "1F4E78", "FFF2CC", "F2F2F2"
    white_side = Side(style="thin", color="FFFFFF")
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(right=white_side)
    for row in ws.iter_rows(min_row=2):
        for index, cell in enumerate(row, 1):
            cell.font = Font(name="Aptos", size=10, color="222222")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor=amber if index in editable_indices else gray)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "D2" if editable_indices else "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 24
    for index, column in enumerate(ws.columns, 1):
        values = [str(cell.value or "") for cell in list(column)[:80]]
        width = min(48, max(10, max((len(value) for value in values), default=0) + 2))
        if editable_indices and index in (1, 2, 3):
            width = (30, 12, 12)[index - 1]
        ws.column_dimensions[column[0].column_letter].width = width
    # Enough height for ordinary records; very long full quotes remain in the cell
    # and can be expanded in Excel without changing their contents.
    for row in ws.iter_rows(min_row=2):
        lines = max((sum(max(1, (len(part) * 1.6) // max(10, ws.column_dimensions[cell.column_letter].width) + 1)
                         for part in str(cell.value or "").split("\n")) for cell in row), default=1)
        ws.row_dimensions[row[0].row].height = min(180, max(28, lines * 15))


def export_workbook(db_path: Path, output: Path) -> dict[str, Any]:
    db_path, output = Path(db_path).resolve(), Path(output).resolve()
    check_output(output, ".xlsx")
    conn = connect_readonly(db_path)
    try:
        conn.execute("BEGIN")
        check_integrity(conn)
        base_hash = state_hash(conn)
        wb = Workbook()
        guide = wb.active
        guide.title = "说明"
        guide.sheet_view.showGridLines = False
        guide["A2"] = "安全隐患知识母库 Excel 编辑快照"
        guide["A2"].font = Font(name="Aptos", size=15, bold=True, color="1F1F1F")
        guide["A4"] = "使用规则"
        guide["A4"].font = Font(name="Aptos", size=11, bold=True, color="1F4E78")
        rules = [
            "本工作簿不是独立母库；每次必须从最新 SQLite 母库导出。",
            "只在黄色列修改，并在该行“操作”填写 update；灰色列禁止修改。",
            "先运行 propose 生成差异提案并审阅，再运行 apply 原子提交。",
            "内容变化增加 revision 并回到待核验；已失效或关闭记录保留关闭状态，旧核验记录保留。",
            "exchange-v1 只更新已有核心记录；新增、删除、合并、标签和别名编辑留给后续版本。",
            "工作簿中的公式一律拒绝，待核验内容仍不能进入正式网站发布数据。",
            "长条款保留完整原文；可在 Excel 展开行高或编辑栏阅读。文本列保持文本格式。",
        ]
        for row, rule in enumerate(rules, 5):
            cell = guide.cell(row, 1, f"{row - 4}. {rule}")
            cell.font = Font(name="Aptos", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            guide.row_dimensions[row].height = 32
        metadata = (("交换格式版本", FORMAT_VERSION), ("母库状态指纹", base_hash),
                    ("导出时间（UTC）", utc_now()), ("母库文件名", db_path.name))
        start = 13
        for offset, (key, value) in enumerate(metadata):
            guide.cell(start + offset, 1, key).font = Font(name="Aptos", size=10, bold=True)
            guide.cell(start + offset, 2, value)
        guide.column_dimensions["A"].width = 92
        guide.column_dimensions["B"].width = 72

        counts: dict[str, int] = {}
        for sheet_name, spec in SHEETS.items():
            ws = wb.create_sheet(sheet_name)
            columns = spec["columns"]
            for index, (_, heading, _, _) in enumerate(columns, 1):
                ws.cell(1, index, heading)
            rows = table_rows(conn, spec["table"])
            counts[sheet_name] = len(rows)
            for row_index, row in enumerate(rows, 2):
                for column_index, (field, _, _, _) in enumerate(columns, 1):
                    write_cell(ws, row_index, column_index, "" if field == "operation" else row[field])
            editable = {i for i, (_, _, can_edit, _) in enumerate(columns, 1) if can_edit}
            _style_table(ws, editable)
            operation_column = next(i for i, (f, _, _, _) in enumerate(columns, 1) if f == "operation")
            dv = DataValidation(type="list", formula1='"update"', allow_blank=True, error="只能留空或填写 update")
            dv.errorTitle = "无效操作"
            dv.showErrorMessage = True
            ws.add_data_validation(dv)
            dv.add(f"{ws.cell(2, operation_column).coordinate}:{ws.cell(max(2, len(rows)+1), operation_column).coordinate}")
            # Highlight rows explicitly marked for import.
            ws.conditional_formatting.add(f"A2:{ws.cell(max(2, len(rows)+1), len(columns)).coordinate}",
                FormulaRule(formula=[f'$C2="update"'], fill=PatternFill("solid", fgColor="D9EAF7")))

        for sheet_name, (table, columns) in READ_ONLY_SHEETS.items():
            ws = wb.create_sheet(sheet_name)
            for index, heading in enumerate(columns, 1):
                ws.cell(1, index, heading)
            rows = table_rows(conn, table)
            for row_index, row in enumerate(rows, 2):
                for column_index, field in enumerate(columns, 1):
                    write_cell(ws, row_index, column_index, row[field])
            _style_table(ws, set())

        wb.calculation.fullCalcOnLoad = False
        blob = BytesIO()
        wb.save(blob)
        wb.close()
        install_output(output, blob.getvalue())
    finally:
        conn.close()
    return {"ok": True, "formatVersion": FORMAT_VERSION, "baseStateHash": base_hash,
            "workbook": str(output), "workbookSha256": sha256_bytes(output.read_bytes()),
            "counts": counts}


def workbook_metadata(wb) -> dict[str, str]:
    ws = wb["说明"] if "说明" in wb.sheetnames else None
    if ws is None:
        raise ValueError("缺少“说明”Sheet")
    values = {}
    for row in ws.iter_rows(min_row=1, max_col=2, values_only=True):
        if row[0] in ("交换格式版本", "母库状态指纹", "导出时间（UTC）", "母库文件名"):
            values[str(row[0])] = str(row[1] or "")
    if values.get("交换格式版本") != FORMAT_VERSION:
        raise ValueError("工作簿交换格式版本不受支持")
    if not values.get("母库状态指纹"):
        raise ValueError("工作簿缺少母库状态指纹")
    return values


def assert_no_formulas(wb) -> None:
    formulas = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" or cell.data_type == "e":
                    formulas.append(f"{ws.title}!{cell.coordinate}")
                    if len(formulas) >= 10:
                        break
    if formulas:
        raise ValueError("工作簿不允许公式或错误单元格: " + ", ".join(formulas))


def check_headers(ws, headings: list[str]) -> None:
    if [ws.cell(1, i).value for i in range(1, len(headings) + 1)] != headings:
        raise ValueError(f"{ws.title}: 列标题或顺序已改变")
    if ws.max_column > len(headings):
        for row in ws.iter_rows(min_col=len(headings) + 1):
            if any(cell.value not in (None, "") for cell in row):
                raise ValueError(f"{ws.title}: 含未知列，拒绝静默忽略")


def check_readonly_sheets(wb, conn: sqlite3.Connection) -> None:
    if set(wb.sheetnames) != {"说明", *SHEETS, *READ_ONLY_SHEETS}:
        raise ValueError("Sheet 集合已改变；不支持新增或删除 Sheet")
    for name, (table, columns) in READ_ONLY_SHEETS.items():
        ws = wb[name]
        check_headers(ws, list(columns))
        expected = Counter(dumps([row[column] if row[column] is not None else "" for column in columns])
                           for row in table_rows(conn, table))
        actual = Counter()
        for cells in ws.iter_rows(min_row=2, max_col=len(columns), values_only=True):
            if all(cell in (None, "") for cell in cells):
                continue
            actual[dumps([cell if cell is not None else "" for cell in cells])] += 1
        if actual != expected:
            raise ValueError(f"{name}: 只读 Sheet 内容已改变")


def make_change(sheet_name: str, before: dict[str, Any], edits: dict[str, Any]) -> dict[str, Any]:
    """One business-rule implementation for both workbook intake and final apply."""
    spec = SHEETS[sheet_name]
    editable = {field: kind for field, _, can_edit, kind in spec["columns"] if can_edit and field != "operation"}
    if not edits or set(edits) - set(editable):
        raise ValueError("变更为空或包含只读字段")
    after = dict(before)
    for field, value in edits.items():
        value = normalize(value, editable[field])
        if value == normalize(before[field], editable[field]):
            raise ValueError(f"{field}: 变更内容与原值相同")
        if editable[field] == "date" and value:
            try:
                if dt.date.fromisoformat(value).isoformat() != value:
                    raise ValueError()
            except ValueError as exc:
                raise ValueError(f"{field}: 日期必须为 YYYY-MM-DD") from exc
        after[field] = value
    for field, value in spec["demote"].items():
        if field in ("status", "review_status") and before[field] in ("已失效", "已废止", "merged", "rejected"):
            continue  # Content editing never reactivates closed records.
        after[field] = value
    after["revision"] = before["revision"] + 1
    for field in spec["required"]:
        if not str(after[field]).strip():
            raise ValueError(f"{field} 不能为空")
    if spec["table"] == "links" and after["priority"] < 0:
        raise ValueError("priority 不能小于 0")
    if spec["table"] == "law_versions":
        if after["validity_status"] not in ("", "待核验", "现行有效", "即将生效", "已废止"):
            raise ValueError("有效性判断非法")
        if {"effective_date", "end_date"} & set(edits) and after["effective_date"] and after["end_date"]:
            if after["end_date"] < after["effective_date"]:
                raise ValueError("失效日期不能早于生效日期")
    if spec["table"] == "laws" and before["identity_status"] != "provisional":
        raise ValueError("已确认或已合并法规身份需要专门的身份审查流程")
    if spec["table"] == "clauses" and "article_path" in edits and before["identity_status"] != "legacy_unreviewed":
        raise ValueError("已确认或冲突中的条款定位需要专门的身份审查流程")
    return {"sheet": sheet_name, "table": spec["table"], "entityType": spec["entity"],
            "id": before["id"], "baseRevision": before["revision"],
            "changedFields": [field for field in editable if field in edits], "before": before, "after": after}


def _read_sheet(wb, sheet_name: str, spec: dict[str, Any], current: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"缺少 Sheet: {sheet_name}")
    ws, columns = wb[sheet_name], spec["columns"]
    check_headers(ws, [column[1] for column in columns])
    workbook_rows: dict[str, tuple[int, dict[str, Any]]] = {}
    for row_index in range(2, ws.max_row + 1):
        raw_id = ws.cell(row_index, 1).value
        if raw_id in (None, "") and all(ws.cell(row_index, i).value in (None, "") for i in range(1, len(columns) + 1)):
            continue
        entity_id = normalize(raw_id, "text")
        if not entity_id:
            raise ValueError(f"{sheet_name}!A{row_index}: ID 不能为空")
        if entity_id in workbook_rows:
            raise ValueError(f"{sheet_name}: 重复 ID {entity_id}")
        row = {}
        for column_index, (field, _, _, kind) in enumerate(columns, 1):
            row[field] = normalize(ws.cell(row_index, column_index).value, kind)
        workbook_rows[entity_id] = (row_index, row)
    if set(workbook_rows) != set(current):
        missing = sorted(set(current) - set(workbook_rows))[:8]
        extra = sorted(set(workbook_rows) - set(current))[:8]
        raise ValueError(f"{sheet_name}: 行集合与导出母库不一致; 缺少={missing}, 新增={extra}")

    changes = []
    for entity_id, (row_index, row) in workbook_rows.items():
        before = current[entity_id]
        normalized_before = {field: normalize(before.get(field), kind) for field, _, _, kind in columns if field != "operation"}
        if row["revision"] != normalized_before["revision"]:
            raise ValueError(f"{sheet_name}!B{row_index}: {entity_id} 基础版本已改变")
        if row["operation"] not in ("", "update"):
            raise ValueError(f"{sheet_name}!C{row_index}: 操作只能留空或填写 update")
        edited_fields = [field for field, _, editable, _ in columns if editable and field != "operation"
                         and row[field] != normalized_before[field]]
        readonly_changes = [field for field, _, editable, _ in columns if not editable
                            and field not in ("revision",) and row[field] != normalized_before[field]]
        if readonly_changes:
            raise ValueError(f"{sheet_name}!{row_index}: {entity_id} 修改了只读列 {readonly_changes}")
        if edited_fields and row["operation"] != "update":
            raise ValueError(f"{sheet_name}!C{row_index}: {entity_id} 有修改但未标记 update")
        if row["operation"] == "update" and not edited_fields:
            raise ValueError(f"{sheet_name}!C{row_index}: {entity_id} 标记 update 但内容没有变化")
        if not edited_fields:
            continue
        changes.append(make_change(sheet_name, before, {field: row[field] for field in edited_fields}))
    return sorted(changes, key=lambda change: change["id"])


def _validate_clause_locators(conn: sqlite3.Connection, changes: list[dict[str, Any]]) -> None:
    affected = {c["id"]: c["after"] for c in changes if c["table"] == "clauses"}
    if not affected:
        return
    rows = table_rows(conn, "clauses")
    groups: dict[tuple[str, str], list[str]] = {}
    for row in rows:
        row = affected.get(row["id"], row)
        groups.setdefault((row["law_version_id"], row["article_path"]), []).append(row["id"])
    moved = {change["id"] for change in changes if change["table"] == "clauses" and "article_path" in change["changedFields"]}
    collisions = [{"lawVersionId": key[0], "articlePath": key[1], "ids": ids}
                  for key, ids in groups.items() if len(ids) > 1 and moved.intersection(ids)]
    if collisions:
        raise ValueError("条款定位发生重复: " + dumps(collisions))


def create_proposal(db_path: Path, workbook_path: Path, output: Path) -> dict[str, Any]:
    db_path, workbook_path, output = map(lambda p: Path(p).resolve(), (db_path, workbook_path, output))
    check_output(output, ".json")
    blob = workbook_path.read_bytes()
    workbook_sha = sha256_bytes(blob)
    wb = load_workbook(BytesIO(blob), data_only=False, read_only=False)
    assert_no_formulas(wb)
    metadata = workbook_metadata(wb)
    conn = connect_readonly(db_path)
    try:
        conn.execute("BEGIN")
        check_integrity(conn)
        current_hash = state_hash(conn)
        if metadata["母库状态指纹"] != current_hash:
            raise ValueError("工作簿对应的母库已变化，请重新导出；拒绝生成过期提案")
        check_readonly_sheets(wb, conn)
        changes = []
        for sheet_name, spec in SHEETS.items():
            rows = {row["id"]: row for row in table_rows(conn, spec["table"])}
            changes.extend(_read_sheet(wb, sheet_name, spec, rows))
        _validate_clause_locators(conn, changes)
    finally:
        conn.close()
        wb.close()
    payload = {"formatVersion": FORMAT_VERSION, "baseStateHash": current_hash,
               "workbookSha256": workbook_sha, "createdAt": utc_now(), "changes": changes}
    proposal_hash = sha256_bytes(dumps(payload).encode("utf-8"))
    proposal = {**payload, "proposalId": "CS_" + proposal_hash[:26], "proposalHash": proposal_hash}
    install_output(output, (json.dumps(proposal, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
    return {"ok": True, "proposal": str(output), "proposalId": proposal["proposalId"],
            "proposalHash": proposal_hash, "changeCount": len(changes),
            "entities": [f'{c["entityType"]}:{c["id"]}' for c in changes]}


def apply_proposal(db_path: Path, proposal_path: Path, actor: str) -> dict[str, Any]:
    if not actor.strip():
        raise ValueError("actor 不能为空")
    db_path, proposal_path = Path(db_path).resolve(), Path(proposal_path).resolve()
    proposal = json.loads(proposal_path.read_text(encoding="utf-8"))
    supplied_hash = proposal.pop("proposalHash", "")
    calculated_hash = sha256_bytes(dumps({k: proposal[k] for k in ("formatVersion", "baseStateHash", "workbookSha256", "createdAt", "changes")}).encode("utf-8"))
    if supplied_hash != calculated_hash or proposal.get("proposalId") != "CS_" + calculated_hash[:26]:
        raise ValueError("提案哈希无效，文件可能被修改")
    if proposal.get("formatVersion") != FORMAT_VERSION:
        raise ValueError("提案格式版本不受支持")
    # mode=rw refuses to create a database at a mistyped path. Schema upgrades,
    # data writes and audit records all belong to the same transaction.
    conn = sqlite3.connect(db_path.as_uri() + "?mode=rw", uri=True)
    conn.execute("PRAGMA foreign_keys=ON")
    try:
        conn.execute("BEGIN IMMEDIATE")
        check_integrity(conn)
        if state_hash(conn) != proposal["baseStateHash"]:
            raise ValueError("母库已变化，拒绝提交过期提案")
        if not isinstance(proposal["changes"], list):
            raise ValueError("changes 必须为数组")
        if not proposal["changes"]:
            conn.rollback()
            return {"ok": True, "status": "no_changes", "changeCount": 0, "proposalId": proposal["proposalId"]}
        seen = set()
        # Reconstruct every change from allowed edits. A checksum detects damage;
        # it is NOT a signature or a replacement for business-rule validation.
        for change in proposal["changes"]:
            if change["sheet"] not in SHEETS:
                raise ValueError("提案包含未知表")
            spec = SHEETS[change["sheet"]]
            key = (spec["table"], change["id"])
            if key in seen:
                raise ValueError("提案包含重复实体")
            seen.add(key)
            current = conn.execute(f'SELECT * FROM "{spec["table"]}" WHERE id=?', (change["id"],)).fetchone()
            if current is None:
                raise ValueError(f'实体不存在: {change["id"]}')
            columns = [row[1] for row in conn.execute(f'PRAGMA table_info("{spec["table"]}")')]
            current_dict = dict(zip(columns, current))
            if current_dict != change["before"] or current_dict["revision"] != change["baseRevision"]:
                raise ValueError(f'实体版本或内容已变化: {change["id"]}')
            if not isinstance(change["changedFields"], list):
                raise ValueError("changedFields 必须为数组")
            edits = {field: change["after"][field] for field in change["changedFields"]}
            canonical = make_change(change["sheet"], current_dict, edits)
            if dumps(change) != dumps(canonical):
                raise ValueError(f'提案不符合字段权限、核验降级或审计规则: {change["id"]}')
        _validate_clause_locators(conn, proposal["changes"])
        # Execute individual DDL statements; executescript would commit first.
        statement = ""
        for line in (ROOT / "source/schemas/master.sql").read_text(encoding="utf-8").splitlines(keepends=True):
            statement += line
            if sqlite3.complete_statement(statement):
                conn.execute(statement)
                statement = ""
        if conn.execute("SELECT 1 FROM change_sets WHERE id=?", (proposal["proposalId"],)).fetchone():
            raise ValueError("该提案已经提交")
        for change in proposal["changes"]:
            spec = SHEETS[change["sheet"]]
            assignments = sorted(key for key in change["after"] if change["after"][key] != change["before"][key])
            sql = ",".join(f'"{field}"=?' for field in assignments)
            conn.execute(f'UPDATE "{spec["table"]}" SET {sql} WHERE id=? AND revision=?',
                         [change["after"][field] for field in assignments] + [change["id"], change["baseRevision"]])
        check_integrity(conn)
        result_hash = state_hash(conn)
        applied_at = utc_now()
        full_proposal = {**proposal, "proposalHash": supplied_hash}
        conn.execute("INSERT INTO change_sets VALUES(?,?,?,?,?,?,?,?,?,?)",
                     (proposal["proposalId"], supplied_hash, proposal["baseStateHash"], result_hash,
                      proposal["workbookSha256"], actor.strip(), "applied", proposal["createdAt"], applied_at, dumps(full_proposal)))
        for ordinal, change in enumerate(proposal["changes"]):
            conn.execute("INSERT INTO change_events VALUES(?,?,?,?,?,?,?,?)",
                         (proposal["proposalId"], ordinal, change["entityType"], change["id"],
                          change["baseRevision"], change["after"]["revision"], dumps(change["before"]), dumps(change["after"])))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return {"ok": True, "proposalId": proposal["proposalId"], "changeCount": len(proposal["changes"]),
            "baseStateHash": proposal["baseStateHash"], "resultStateHash": result_hash,
            "actor": actor.strip(), "appliedAt": applied_at}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, default=ROOT / "source/master/safety.sqlite3")
    sub = parser.add_subparsers(dest="command", required=True)
    exp = sub.add_parser("export", help="export the latest editable workbook")
    exp.add_argument("--output", type=Path, required=True)
    prop = sub.add_parser("propose", help="validate a workbook and create a JSON change proposal")
    prop.add_argument("--workbook", type=Path, required=True)
    prop.add_argument("--output", type=Path, required=True)
    app = sub.add_parser("apply", help="atomically apply one reviewed proposal")
    app.add_argument("--proposal", type=Path, required=True)
    app.add_argument("--actor", required=True)
    args = parser.parse_args()
    try:
        if args.command == "export":
            report = export_workbook(args.db, args.output)
        elif args.command == "propose":
            report = create_proposal(args.db, args.workbook, args.output)
        else:
            report = apply_proposal(args.db, args.proposal, args.actor)
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False, indent=2))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

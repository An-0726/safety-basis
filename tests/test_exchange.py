"""Behavioral checks for the private master editing boundary, using real legacy data."""
import json
from pathlib import Path
import shutil
import sqlite3
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "tools/pipeline"))
import exchange
import master


class ExchangeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.fixture_dir = tempfile.TemporaryDirectory()
        cls.fixture = Path(cls.fixture_dir.name)
        master.migrate(REPO, cls.fixture / "master.sqlite3", cls.fixture / "archive")
        exchange.export_workbook(cls.fixture / "master.sqlite3", cls.fixture / "base.xlsx")

    @classmethod
    def tearDownClass(cls):
        cls.fixture_dir.cleanup()

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.db = self.root / "master.sqlite3"
        self.book = self.root / "edit.xlsx"
        self.proposal = self.root / "proposal.json"
        shutil.copyfile(self.fixture / "master.sqlite3", self.db)
        shutil.copyfile(self.fixture / "base.xlsx", self.book)

    def tearDown(self):
        self.tmp.cleanup()

    def edit(self, changes):
        wb = load_workbook(self.book)
        for sheet, cell, value in changes:
            wb[sheet][cell] = value
        wb.save(self.book)
        wb.close()

    def prepare(self):
        self.edit([("隐患库", "C2", "update"), ("隐患库", "D2", "测试修订后的隐患")])
        return exchange.create_proposal(self.db, self.book, self.proposal)

    def fetch(self, sql):
        with sqlite3.connect(self.db) as conn:
            result = conn.execute(sql).fetchall()
        conn.close()
        return result

    def modify_db(self, sql):
        conn = sqlite3.connect(self.db)
        try:
            with conn:
                conn.execute(sql)
        finally:
            conn.close()

    def rehash(self, payload):
        body = {key: payload[key] for key in ("formatVersion", "baseStateHash", "workbookSha256", "createdAt", "changes")}
        sha = exchange.sha256_bytes(master.dumps(body).encode("utf-8"))
        payload.update(proposalId="CS_" + sha[:26], proposalHash=sha)
        self.proposal.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    def test_export_and_reordered_roundtrip_preserve_database_bytes(self):
        before = self.db.read_bytes()
        wb = load_workbook(self.book)
        for ws in wb.worksheets:
            if ws.title == "说明" or ws.max_row < 3:
                continue
            rows = list(ws.iter_rows(min_row=2, values_only=True))[::-1]
            for r, row in enumerate(rows, 2):
                for c, value in enumerate(row, 1):
                    ws.cell(r, c).value = value
        wb.save(self.book)
        wb.close()
        self.assertEqual(exchange.create_proposal(self.db, self.book, self.proposal)["changeCount"], 0)
        self.assertEqual(exchange.apply_proposal(self.db, self.proposal, "test")["status"], "no_changes")
        self.assertEqual(before, self.db.read_bytes())

    def test_apply_changes_revision_demotes_and_preserves_evidence(self):
        evidence = self.fetch("SELECT * FROM verification")
        self.prepare()
        result = exchange.apply_proposal(self.db, self.proposal, "test reviewer")
        self.assertEqual(result["changeCount"], 1)
        self.assertEqual(self.fetch("SELECT title,revision,status,checked FROM hazards WHERE id='H001'")[0],
                         ("测试修订后的隐患", 2, "待核验", ""))
        self.assertEqual(evidence, self.fetch("SELECT * FROM verification"))
        event = self.fetch("SELECT before_json,after_json FROM change_events")[0]
        self.assertEqual(json.loads(event[0])["revision"], 1)
        self.assertEqual(json.loads(event[1])["revision"], 2)
        with self.assertRaises(ValueError):
            exchange.apply_proposal(self.db, self.proposal, "test reviewer")
        self.assertEqual(self.fetch("SELECT COUNT(*) FROM change_sets")[0][0], 1)

    def test_all_five_entities_edit_in_one_transaction(self):
        self.edit([("隐患库", "C2", "update"), ("隐患库", "D2", "现场条件补充"),
                   ("法规库", "C2", "update"), ("法规库", "E2", "待核验发布机关"),
                   ("法规版本", "C2", "update"), ("法规版本", "I2", "待核验范围"),
                   ("条款库", "C2", "update"), ("条款库", "G2", "https://example.gov.cn/evidence"),
                   ("依据关联", "C2", "update"), ("依据关联", "H2", "待核验适用条件")])
        self.assertEqual(exchange.create_proposal(self.db, self.book, self.proposal)["changeCount"], 5)
        self.assertEqual(exchange.apply_proposal(self.db, self.proposal, "integration")["changeCount"], 5)
        self.assertEqual(self.fetch("SELECT COUNT(*) FROM change_events")[0][0], 5)
        for table in ("hazards", "laws", "law_versions", "clauses", "links"):
            status = "review_status" if table == "law_versions" else "status"
            self.assertEqual(self.fetch(f"SELECT {status} FROM {table} WHERE revision=2"), [("待核验",)])

    def test_stale_workbook_and_proposal_rejected(self):
        self.prepare()
        self.modify_db("UPDATE hazards SET note='another edit',revision=revision+1 WHERE id='H002'")
        before = self.db.read_bytes()
        with self.assertRaisesRegex(ValueError, "母库.*变化"):
            exchange.create_proposal(self.db, self.book, self.root / "new.json")
        with self.assertRaisesRegex(ValueError, "母库.*变化"):
            exchange.apply_proposal(self.db, self.proposal, "test")
        self.assertEqual(before, self.db.read_bytes())

    def test_unmarked_empty_invalid_and_readonly_edits_rejected(self):
        scenarios = [
            [("隐患库", "D2", "unmarked")],
            [("隐患库", "C2", "update")],
            [("隐患库", "C2", "delete")],
            [("隐患库", "K2", "已核验")],
            [("隐患库", "B2", 99)],
            [("依据关联", "D2", "missing")],
            [("隐患库", "C2", "update"), ("隐患库", "D2", "")],
            [("依据关联", "C2", "update"), ("依据关联", "G2", None)],
            [("依据关联", "C2", "update"), ("依据关联", "G2", -1)],
            [("法规版本", "C2", "update"), ("法规版本", "J2", "2026-99-99")],
            [("法规版本", "C2", "update"), ("法规版本", "L2", "随便填的状态")],
        ]
        # H001 may already be verified; use a different value for its readonly status.
        scenarios[3] = [("隐患库", "K2", "私自改状态")]
        for changes in scenarios:
            with self.subTest(changes=changes):
                shutil.copyfile(self.fixture / "base.xlsx", self.book)
                self.edit(changes)
                with self.assertRaises(ValueError):
                    exchange.create_proposal(self.db, self.book, self.proposal)
                self.assertFalse(self.proposal.exists())

    def test_formula_in_core_or_readonly_sheet_rejected(self):
        for name in ("隐患库", "原始来源"):
            with self.subTest(sheet=name):
                shutil.copyfile(self.fixture / "base.xlsx", self.book)
                self.edit([(name, "D2", "=1+1")])
                with self.assertRaisesRegex(ValueError, "公式"):
                    exchange.create_proposal(self.db, self.book, self.proposal)

    def test_deleted_duplicate_new_rows_and_sheet_structure_rejected(self):
        for kind in ("deleted", "duplicate", "new", "missing_sheet", "new_sheet", "extra_column", "readonly_edit", "readonly_deleted"):
            with self.subTest(kind=kind):
                shutil.copyfile(self.fixture / "base.xlsx", self.book)
                wb = load_workbook(self.book)
                ws = wb["隐患库"]
                if kind == "deleted": ws.delete_rows(2)
                if kind == "duplicate": ws["A3"] = ws["A2"].value
                if kind == "new": ws.cell(ws.max_row + 1, 1, "H999")
                if kind == "missing_sheet": del wb["核验记录"]
                if kind == "new_sheet": wb.create_sheet("误建表")
                if kind == "extra_column": ws["Z2"] = "不能丢弃"
                if kind == "readonly_edit": wb["法规别名"]["B2"] = "不能忽略"
                if kind == "readonly_deleted": wb["原始来源"].delete_rows(2)
                wb.save(self.book)
                wb.close()
                with self.assertRaises(ValueError):
                    exchange.create_proposal(self.db, self.book, self.proposal)

    def test_rehashed_invalid_proposals_cannot_bypass_rules(self):
        self.prepare()
        original = json.loads(self.proposal.read_text(encoding="utf-8"))
        for kind in ("keep_verified", "checked", "empty_title", "revision", "readonly", "entity_type", "drop_field", "duplicate", "negative_priority"):
            with self.subTest(kind=kind):
                data = json.loads(json.dumps(original))
                change = data["changes"][0]
                if kind == "keep_verified": change["after"]["status"] = "已核验"
                if kind == "checked": change["after"]["checked"] = "2026-09-08"
                if kind == "empty_title": change["after"]["title"] = ""
                if kind == "revision": change["after"]["revision"] = 10
                if kind == "readonly": change["after"]["merged_into"] = "H002"
                if kind == "entity_type": change["entityType"] = "law"
                if kind == "drop_field": del change["after"]["legacy_payload"]
                if kind == "duplicate": data["changes"].append(change.copy())
                if kind == "negative_priority": change["changedFields"].append("priority"); change["after"]["priority"] = -1
                self.rehash(data)
                before = self.db.read_bytes()
                with self.assertRaises(ValueError):
                    exchange.apply_proposal(self.db, self.proposal, "test")
                self.assertEqual(before, self.db.read_bytes())

    def test_plain_tampering_and_empty_actor_rejected(self):
        self.prepare()
        with self.assertRaises(ValueError): exchange.apply_proposal(self.db, self.proposal, " ")
        data = json.loads(self.proposal.read_text(encoding="utf-8"))
        data["changes"][0]["after"]["title"] = "改提案不重算校验和"
        self.proposal.write_text(json.dumps(data), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "哈希"):
            exchange.apply_proposal(self.db, self.proposal, "test")

    def test_late_failure_rolls_back_entity_and_schema_upgrade(self):
        self.prepare()
        self.modify_db("DROP TABLE change_events")
        self.modify_db("DROP TABLE change_sets")
        real_hash = exchange.state_hash
        calls = 0
        def fail_after_update(conn):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise RuntimeError("forced failure after data and schema writes")
            return real_hash(conn)
        before = self.db.read_bytes()
        with patch.object(exchange, "state_hash", side_effect=fail_after_update):
            with self.assertRaises(RuntimeError):
                exchange.apply_proposal(self.db, self.proposal, "test")
        self.assertEqual(before, self.db.read_bytes())
        self.assertEqual(self.fetch("SELECT name FROM sqlite_master WHERE name='change_sets'"), [])

    def test_old_master_schema_upgrades_only_on_success(self):
        self.prepare()
        self.modify_db("DROP TABLE change_events")
        self.modify_db("DROP TABLE change_sets")
        exchange.apply_proposal(self.db, self.proposal, "test")
        self.assertEqual(self.fetch("SELECT COUNT(*) FROM change_events"), [(1,)])
        with self.assertRaises(sqlite3.IntegrityError):
            self.modify_db("DELETE FROM change_events")
        with self.assertRaises(sqlite3.IntegrityError):
            self.modify_db("UPDATE change_sets SET actor='rewrite history'")

    def test_output_never_overwrites_inputs_or_existing_work(self):
        before_db, before_book = self.db.read_bytes(), self.book.read_bytes()
        for output in (self.db, self.book):
            with self.subTest(output=output):
                with self.assertRaises(ValueError): exchange.export_workbook(self.db, output)
                with self.assertRaises(ValueError): exchange.create_proposal(self.db, self.book, output)
        self.assertEqual(before_db, self.db.read_bytes())
        self.assertEqual(before_book, self.book.read_bytes())

    def test_missing_database_is_not_created(self):
        self.prepare()
        missing = self.root / "missing.sqlite3"
        with self.assertRaises(sqlite3.OperationalError):
            exchange.apply_proposal(missing, self.proposal, "test")
        self.assertFalse(missing.exists())

    def test_whitespace_and_formula_like_literal_text_are_preserved(self):
        self.modify_db("UPDATE hazards SET note='  =literally text  ' WHERE id='H001'")
        self.book = self.root / "literal.xlsx"
        exchange.export_workbook(self.db, self.book)
        self.assertEqual(exchange.create_proposal(self.db, self.book, self.proposal)["changeCount"], 0)

    def test_excel_long_text_limit_fails_without_partial_output(self):
        self.modify_db("UPDATE clauses SET quote=replace(hex(zeroblob(20000)),'0','文') WHERE id='C001'")
        target = self.root / "long.xlsx"
        with self.assertRaisesRegex(ValueError, "拒绝截断"):
            exchange.export_workbook(self.db, target)
        self.assertFalse(target.exists())

    def test_closed_hazard_is_not_reactivated(self):
        self.modify_db("UPDATE hazards SET status='已失效' WHERE id='H001'")
        self.book = self.root / "closed.xlsx"
        exchange.export_workbook(self.db, self.book)
        self.prepare()
        exchange.apply_proposal(self.db, self.proposal, "test")
        self.assertEqual(self.fetch("SELECT status,revision FROM hazards WHERE id='H001'"), [("已失效", 2)])

    def test_new_locator_collision_is_rejected(self):
        wb = load_workbook(self.book)
        ws = wb["条款库"]
        first = ws["D2"].value
        peer = next(r for r in range(3, ws.max_row + 1) if ws.cell(r, 4).value == first and ws.cell(r, 5).value != ws["E2"].value)
        ws["C2"] = "update"
        ws["E2"] = ws.cell(peer, 5).value
        wb.save(self.book)
        wb.close()
        with self.assertRaisesRegex(ValueError, "定位.*重复"):
            exchange.create_proposal(self.db, self.book, self.proposal)

    def test_read_transaction_keeps_one_snapshot_during_concurrent_write(self):
        self.modify_db("PRAGMA journal_mode=WAL")
        original = exchange.table_rows
        raced = False
        def race(conn, table):
            nonlocal raced
            if table == "hazards" and not raced:
                raced = True
                self.modify_db("UPDATE hazards SET note='concurrent',revision=revision+1 WHERE id='H001'")
            return original(conn, table)
        target = self.root / "snapshot.xlsx"
        with patch.object(exchange, "table_rows", side_effect=race):
            exchange.export_workbook(self.db, target)
        wb = load_workbook(target)
        self.assertEqual(wb["隐患库"]["B2"].value, 1)
        self.assertNotEqual(wb["隐患库"]["I2"].value, "concurrent")
        wb.close()
        with self.assertRaisesRegex(ValueError, "母库.*变化"):
            exchange.create_proposal(self.db, target, self.proposal)


if __name__ == "__main__":
    unittest.main()

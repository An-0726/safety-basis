# -*- coding: utf-8 -*-
"""把用户审查工作簿落成处置：阻断条目删除（git 留底），保留条目不动。
裁决与修改建议全部存 tmp/review_verdicts.json，作为返工阶段的依据。"""
import glob
import io
import json
import os
from collections import Counter

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
KNOW = os.path.join(ROOT, "knowledge")
XLSX = r"D:\weixin\data\xwechat_files\wxid_vgdn1gpt0yi722_4e80\msg\file\2026-09\隐患库上线数据详细审查_20260913.xlsx"

wb = load_workbook(XLSX, read_only=True)
ws = wb["全量逐条审查"]
rows = ws.iter_rows(values_only=True)
head = list(next(rows))

verdicts = {}
for row in rows:
    if not row[1]:
        continue
    d = {head[i]: (str(x) if x is not None else "") for i, x in enumerate(row)}
    verdicts[d["隐患ID"]] = d

cnt = Counter(v["审查结论"] for v in verdicts.values())
print("审查结论分布:", dict(cnt))

io.open(os.path.join(ROOT, "tmp", "review_verdicts.json"), "w", encoding="utf-8", newline="\n").write(
    json.dumps(verdicts, ensure_ascii=False, indent=1))

KEEP_PREFIX = ("修改后可上线", "基本可用")
removed = {"hazards": 0, "links": 0, "rev_links": 0, "rev_hazards": 0}
missing = []
for hid, v in verdicts.items():
    if v["审查结论"].startswith(KEEP_PREFIX):
        continue
    hp = os.path.join(KNOW, "hazards", hid + ".json")
    if not os.path.exists(hp):
        missing.append(hid)
        continue
    for lp in glob.glob(os.path.join(KNOW, "links", "*.json")):
        l = json.load(io.open(lp, encoding="utf-8"))
        if l.get("hazardId") == hid and l.get("lifecycle") == "active":
            os.remove(lp)
            removed["links"] += 1
            rp = os.path.join(KNOW, "reviews", "links", os.path.basename(lp))
            if os.path.exists(rp):
                os.remove(rp)
                removed["rev_links"] += 1
    os.remove(hp)
    removed["hazards"] += 1
    rp = os.path.join(KNOW, "reviews", "hazards", hid + ".json")
    if os.path.exists(rp):
        os.remove(rp)
        removed["rev_hazards"] += 1

print("removed:", removed)
print("no knowledge record for:", len(missing), missing[:5])
